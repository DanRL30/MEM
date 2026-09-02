"""Genera la plantilla canónica de inputs de un caso.

La plantilla no es un archivo fijo: se construye a partir de las unidades
productivas que el caso declare, porque el catálogo de inputs demostró que las
hojas del modelo repiten el mismo bloque de conceptos por unidad. Un proyecto
con una mina y un metal recibe la misma estructura que uno con cuatro unidades
y tres metales, instanciada menos veces.

Se genera desde `docs/modelo-economico/catalogo-inputs.md`, del que este script
es la contraparte ejecutable: los conceptos de aquí y los del catálogo son los
mismos, y si divergen, manda el catálogo y este archivo se corrige.

**La produccion sale en una pestana por proyecto**, con los sub-bloques `Mina` y
`Planta` que usa el libro: de donde sale el mineral y por que proceso pasa. Cada
corriente de tonelaje lleva debajo la ley de cada metal que transporta.

Uso:
    python scripts/generar_plantilla_inputs.py --salida produccion.xlsx \
        --unidad "Proyecto X:mina:Sn,Cu" \
        --unidad "San Rafael:mina:Sn:preconcentracion,concentradora" \
        --unidad "B2:mina:Sn:concentradora:relave" \
        --unidad "Pisco:fundicion:Sn" \
        --primer-ano 2027 --anos 20

El cuarto campo declara las etapas de la planta y el quinto el origen del
mineral; los dos son opcionales. Sin etapas, una mina se describe con su
concentradora y nada mas, que es el caso general: pedir datos de preconcentracion
a un proyecto que no la tiene es pedir un dato que no existe.

El origen distingue una mina de una relavera. **Una relavera no es una etapa de
tratamiento: es una mina cuyo mineral sale de un deposito de relaves ya cerrado**,
con su tonelaje extraido y su ley igual que cualquier otra, y desde ahi la misma
cadena.

Si el caso declara una fundicion, todas las minas le entregan su concentrado. Un
proyecto que venda directo simplemente no declara ninguna.

Las celdas verdes son **corroborables**: se llenan igual que las demas, y ademas
el sistema las recalcula y avisa si el dato cargado no cuadra con la cadena.

La plantilla sale **vacía**. Llenarla con datos de un caso real la convierte en
información confidencial de MINSUR: no vuelve al repositorio.
"""

from __future__ import annotations

import argparse
import sys
from collections.abc import Mapping
from dataclasses import dataclass, field, replace
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover - entorno sin sincronizar
    print("Falta openpyxl. Sincroniza el entorno con: uv sync --all-packages")
    raise SystemExit(1) from None

TIPOS_DE_UNIDAD = ("mina", "fundicion")
ORIGENES = ("yacimiento", "relave")
ETAPAS = ("preconcentracion", "concentradora")
ETAPAS_POR_DEFECTO = ("concentradora",)
METALES = ("Sn", "Cu", "Ag")

# Corrientes de mineral de una unidad, en el orden del libro. Cada una lleva su
# tonelaje y, debajo, la ley de cada metal que transporta. El libro las escribe
# asi, en pares contiguos, y rotula todas las leyes igual: lo unico que las
# distingue es esa vecindad. Aqui cada ley se nombra por la corriente que
# describe, que es lo que permite leerlas sin depender de la posicion.
#
#   (tonelaje, etapa que lo condiciona, como se nombra su ley, si se corrobora)
CORRIENTES_DE_MINERAL = [
    ("Mineral extraido", None, "del mineral extraido", False),
    (
        "Mineral tratado en preconcentracion",
        "preconcentracion",
        "de entrada a preconcentracion",
        False,
    ),
    ("Mineral preconcentrado a concentradora", "preconcentracion", "del preconcentrado", False),
    ("Mineral directo a planta concentradora", "concentradora", "del mineral directo", False),
    ("Mineral tratado total en concentradora", "concentradora", "del tratado total", True),
    ("Mineral tratado total para cash cost", None, "del tratado para cash cost", False),
]
# Lo que la planta produce para un metal con concentrado propio. El orden es el
# del libro, que es el que Finanzas reconoce al revisar la plantilla.
CONCENTRADO_POR_METAL = [
    ("Toneladas finas de {metal}", "tmf", True),
    ("Ley de {metal} en el concentrado", "%", False),
    ("Recuperacion de {metal}", "%", False),
    ("Produccion de concentrado de {metal}", "t", True),
]
# Filas de la unidad de fundicion. La capacidad es la regla 002: el libro la
# escribia dentro de la formula y Finanzas confirmo el 01/09/2026 que es dato
# editable, asi que es una fila mas.
COMPLEJO = [
    ("Toneladas alimentadas mas escoria", "t", False),
    ("Capacidad maxima de tratamiento", "t", False),
    ("Concentrado excedente", "t", False),
]
COMPLEJO_POR_METAL = [
    ("Ley promedio de alimentacion de {metal}", "%", True),
    ("Produccion de metal refinado de {metal}", "tmf", False),
]
OPEX_POR_UNIDAD = [
    ("Exploraciones", None),
    ("Geologia", None),
    ("Mina", None),
    ("Planta de preconcentracion", "preconcentracion"),
    ("Planta concentradora", "concentradora"),
    ("Fundicion", "fundicion"),
    ("Refineria", "fundicion"),
    ("Planta de subproductos", "fundicion"),
    ("Mantenimiento", None),
    ("Energia", None),
    ("Linea de transmision", None),
    ("Agua potable", None),
    ("Planilla", None),
    ("Apoyo", None),
    ("Gestion social", None),
    ("Predios, servidumbres y usufructos", None),
    ("Estudios y optimizaciones", None),
    ("Relavera", "relavera"),
]
CAPEX_ETAPAS = ["Capex inicial", "Sostenimiento", "Cierre de mina", "Otros"]
CAPEX_NATURALEZAS = [
    "No depreciable",
    "Maquinaria, equipos y vehiculos",
    "Instalaciones y equipos diversos y de comunicaciones",
    "Edificaciones y construcciones",
]
DATOS_COMUNES = [
    ("Dias de working capital", "dias"),
    ("Tratamiento del IGV en working capital", "si / no"),
    ("Perdidas tributarias arrastradas", "US$"),
    ("Costos hundidos excluidos del flujo", "US$"),
    ("Inversion social", "US$/ano"),
    ("Gastos administrativos", "US$/ano"),
    ("Otros gastos operativos", "US$/ano"),
    ("Servidumbres y usufructos", "US$/ano"),
    ("Estudios", "US$/ano"),
    ("Exploraciones no atribuibles a una unidad", "US$/ano"),
]
ESCENARIOS_DE_PRECIO = ["Base", "Alto", "Bajo"]

TITULO = Font(bold=True, size=12)
CABECERA = Font(bold=True)
SECCION = PatternFill("solid", fgColor="DDDDDD")
ENTRADA = PatternFill("solid", fgColor="FFF6CC")
CORROBORABLE = PatternFill("solid", fgColor="E6F2E6")


def _leer_metales(nombre: str, texto: str) -> tuple[tuple[str, ...], dict[str, str]]:
    """Interpreta `Sn,Cu,Ag@Cu`: la plata viaja en el concentrado de cobre."""
    lista: list[str] = []
    portador: dict[str, str] = {}
    for bruto in texto.split(","):
        pieza = bruto.strip()
        if not pieza:
            continue
        metal, _, dentro_de = pieza.partition("@")
        metal = metal.strip()
        if metal not in METALES:
            raise ValueError(
                f"Metal {metal!r} fuera del alcance en {nombre!r}. Use {', '.join(METALES)}"
            )
        lista.append(metal)
        if dentro_de:
            portador[metal] = dentro_de.strip()
    if not lista:
        raise ValueError(f"La unidad {nombre!r} no declara metales")
    huerfanos = [m for m, p in portador.items() if p not in lista]
    if huerfanos:
        raise ValueError(
            f"En {nombre!r}, {', '.join(huerfanos)} viaja en el concentrado de un metal que la "
            "unidad no declara."
        )
    return tuple(lista), portador


@dataclass(frozen=True)
class Unidad:
    """Una unidad productiva declarada por el caso.

    El tipo dice si saca mineral o si recibe concentrado, y el origen dice de
    donde lo saca. Una relavera es una mina cuyo origen es un deposito de
    relaves ya cerrado: tiene mineral extraido y ley igual que cualquier otra,
    y de ahi sigue la misma cadena. Tratarla como una etapa de tratamiento
    aparte, que es lo que esta plantilla hacia antes, no corresponde a ningun
    bloque del libro.
    """

    nombre: str
    tipo: str
    metales: tuple[str, ...]
    etapas: tuple[str, ...]
    origen: str = "yacimiento"
    entrega_a: str | None = None
    portador: Mapping[str, str] = field(default_factory=dict)
    """Metal que viaja dentro del concentrado de otro, y en cuál.

    En el libro la plata no tiene concentrado propio: va dentro del de cobre, y
    lo único que se declara de ella es su ley en ese concentrado. Pedirle a un
    proyecto la producción de un concentrado de plata sería pedir un dato que no
    existe.
    """

    def con_concentrado_propio(self) -> tuple[str, ...]:
        return tuple(m for m in self.metales if m not in self.portador)

    @classmethod
    def desde_texto(cls, texto: str) -> Unidad:
        partes = texto.split(":")
        if not 3 <= len(partes) <= 5:
            raise ValueError(
                f"Formato esperado nombre:tipo:metales[:etapas][:origen], recibido {texto!r}"
            )
        nombre, tipo, metales = (p.strip() for p in partes[:3])
        if tipo not in TIPOS_DE_UNIDAD:
            raise ValueError(f"Tipo {tipo!r} desconocido. Use uno de {', '.join(TIPOS_DE_UNIDAD)}")
        lista, portador = _leer_metales(nombre, metales)

        if len(partes) >= 4 and partes[3].strip():
            etapas = tuple(e.strip() for e in partes[3].split(",") if e.strip())
            desconocidas = [e for e in etapas if e not in ETAPAS]
            if desconocidas:
                raise ValueError(
                    f"Etapa(s) {', '.join(desconocidas)} desconocida(s) en {nombre!r}. "
                    f"Use {', '.join(ETAPAS)}"
                )
        elif tipo == "mina":
            etapas = ETAPAS_POR_DEFECTO
        else:
            etapas = ()

        origen = partes[4].strip() if len(partes) == 5 and partes[4].strip() else "yacimiento"
        if origen not in ORIGENES:
            raise ValueError(
                f"Origen {origen!r} desconocido en {nombre!r}. Use {', '.join(ORIGENES)}"
            )
        return cls(nombre, tipo, lista, etapas, origen, portador=portador)

    @property
    def es_fundicion(self) -> bool:
        return self.tipo == "fundicion"

    def aplica(self, etapa: str | None) -> bool:
        """Si la fila condicionada por `etapa` corresponde a esta unidad."""
        if etapa is None:
            return True
        if etapa == "fundicion":
            return self.es_fundicion
        if etapa == "relavera":
            return self.origen == "relave"
        return etapa in self.etapas


def encabezar(hoja: Worksheet, titulo: str, primer_ano: int, anos: int) -> None:
    """Escribe el título y la fila de años, y fija los paneles."""
    hoja["A1"] = titulo
    hoja["A1"].font = TITULO
    hoja["A3"] = "Concepto"
    hoja["B3"] = "Unidad"
    hoja["A3"].font = CABECERA
    hoja["B3"].font = CABECERA
    for i in range(anos):
        celda = hoja.cell(row=3, column=3 + i, value=primer_ano + i)
        celda.font = CABECERA
        celda.alignment = Alignment(horizontal="center")
    hoja.column_dimensions["A"].width = 52
    hoja.column_dimensions["B"].width = 12
    for i in range(anos):
        hoja.column_dimensions[get_column_letter(3 + i)].width = 11
    hoja.freeze_panes = "C4"


def fila_de_seccion(hoja: Worksheet, fila: int, texto: str, anos: int) -> int:
    celda = hoja.cell(row=fila, column=1, value=texto)
    celda.font = CABECERA
    for col in range(1, 3 + anos):
        hoja.cell(row=fila, column=col).fill = SECCION
    return fila + 1


def fila_de_entrada(
    hoja: Worksheet,
    fila: int,
    concepto: str,
    unidad: str,
    anos: int,
    corroborable: bool = False,
) -> int:
    """Escribe una fila de datos.

    Las corroborables se pintan distinto: son las que el sistema recalcula para
    avisar si el dato cargado no cuadra. Se llenan igual que las demas, porque
    el recalculo audita el dato y no lo sustituye.
    """
    hoja.cell(row=fila, column=1, value=concepto)
    hoja.cell(row=fila, column=2, value=unidad)
    relleno = CORROBORABLE if corroborable else ENTRADA
    for col in range(3, 3 + anos):
        hoja.cell(row=fila, column=col).fill = relleno
    return fila + 1


def validacion_numerica(hoja: Worksheet, filas: list[int], anos: int) -> None:
    """Impide texto en las celdas de datos, que es el error de carga más común."""
    if not filas:
        return
    rangos = " ".join(f"C{f}:{get_column_letter(2 + anos)}{f}" for f in filas)
    validacion = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="-1000000000",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Valor no numerico",
        error="Esta celda espera un numero. Deje la celda vacia si el dato no aplica.",
    )
    hoja.add_data_validation(validacion)
    validacion.sqref = rangos


def hoja_caso(libro: Workbook, unidades: list[Unidad], primer_ano: int, anos: int) -> None:
    hoja = libro.active
    hoja.title = "Caso"
    hoja["A1"] = "Identificacion del caso"
    hoja["A1"].font = TITULO
    campos = [
        ("Nombre del caso", ""),
        ("Categoria", ""),
        ("Responsable", ""),
        ("Primer ano del horizonte", primer_ano),
        ("Numero de anos", anos),
        ("Ano de valuacion", ""),
        ("Version de datos maestros", ""),
        ("Escenario de precios", ESCENARIOS_DE_PRECIO[0]),
    ]
    for i, (etiqueta, valor) in enumerate(campos, start=3):
        hoja.cell(row=i, column=1, value=etiqueta).font = CABECERA
        celda = hoja.cell(row=i, column=2, value=valor)
        celda.fill = ENTRADA

    fila = len(campos) + 5
    hoja.cell(row=fila, column=1, value="Unidades productivas declaradas").font = TITULO
    fila += 2
    # Las etapas y el origen viajan en la tabla y no en la hoja de
    # instrucciones: son los que deciden que filas tiene cada pestana, y si el
    # lector no los ve, de una plantilla llena no se puede regenerar la misma
    # plantilla.
    for columna, titulo in enumerate(
        ("Unidad", "Tipo", "Metales", "Origen", "Etapas", "Entrega a"), start=1
    ):
        hoja.cell(row=fila, column=columna, value=titulo).font = CABECERA
    fila += 1
    for unidad in unidades:
        hoja.cell(row=fila, column=1, value=unidad.nombre)
        hoja.cell(row=fila, column=2, value=unidad.tipo)
        hoja.cell(row=fila, column=3, value=", ".join(unidad.metales))
        hoja.cell(row=fila, column=4, value=unidad.origen)
        hoja.cell(row=fila, column=5, value=", ".join(unidad.etapas))
        hoja.cell(row=fila, column=6, value=unidad.entrega_a or "")
        fila += 1

    fila += 2
    hoja.cell(row=fila, column=1, value="Datos comunes al caso").font = TITULO
    fila += 2
    for concepto, unidad_medida in DATOS_COMUNES:
        hoja.cell(row=fila, column=1, value=concepto)
        hoja.cell(row=fila, column=2, value=unidad_medida)
        hoja.cell(row=fila, column=3).fill = ENTRADA
        fila += 1

    hoja.column_dimensions["A"].width = 52
    hoja.column_dimensions["B"].width = 22
    hoja.column_dimensions["C"].width = 22


def hoja_produccion_de_unidad(
    libro: Workbook, unidad: Unidad, unidades: list[Unidad], primer_ano: int, anos: int
) -> None:
    """Una pestana por proyecto, con los sub-bloques del libro.

    El libro describe cada unidad con dos sub-bloques rotulados `Mina` y
    `Planta`: de donde sale el mineral y por que proceso pasa. Reproducirlos es
    lo que hace que Finanzas reconozca la plantilla al revisarla.
    """
    hoja = libro.create_sheet(_nombre_de_hoja(unidad.nombre))
    encabezar(hoja, f"Produccion de {unidad.nombre}", primer_ano, anos)
    hoja["A2"] = _descripcion(unidad)
    fila, numericas = 5, []

    if unidad.es_fundicion:
        fila, numericas = _bloque_de_complejo(hoja, unidad, unidades, fila, numericas, anos)
    else:
        fila, numericas = _bloque_de_mina(hoja, unidad, fila, numericas, anos)
    validacion_numerica(hoja, numericas, anos)


def _bloque_de_mina(
    hoja: Worksheet, unidad: Unidad, fila: int, numericas: list[int], anos: int
) -> tuple[int, list[int]]:
    fila = fila_de_seccion(hoja, fila, "Mina", anos)
    extraido, resto = CORRIENTES_DE_MINERAL[0], CORRIENTES_DE_MINERAL[1:]
    fila, numericas = _corriente(hoja, unidad, extraido, fila, numericas, anos)

    fila = fila_de_seccion(hoja, fila, "Planta", anos)
    for corriente in resto:
        if not unidad.aplica(corriente[1]):
            continue
        fila, numericas = _corriente(hoja, unidad, corriente, fila, numericas, anos)

    for metal in unidad.con_concentrado_propio():
        for plantilla, medida, corroborable in CONCENTRADO_POR_METAL:
            numericas.append(fila)
            fila = fila_de_entrada(
                hoja, fila, plantilla.format(metal=metal), medida, anos, corroborable
            )
        # Los subproductos que viajan en este concentrado no tienen concentrado
        # propio: de ellos solo se declara su ley dentro del que los transporta.
        for subproducto, dentro_de in sorted(unidad.portador.items()):
            if dentro_de != metal:
                continue
            numericas.append(fila)
            fila = fila_de_entrada(
                hoja, fila, f"Ley de {subproducto} en el concentrado de {metal}", "%", anos
            )

    if unidad.entrega_a:
        numericas.append(fila)
        fila = fila_de_entrada(hoja, fila, "Concentrado entregado al complejo", "t", anos, True)
    return fila, numericas


def _bloque_de_complejo(
    hoja: Worksheet,
    unidad: Unidad,
    unidades: list[Unidad],
    fila: int,
    numericas: list[int],
    anos: int,
) -> tuple[int, list[int]]:
    """El complejo no extrae ni trata mineral: recibe concentrado.

    Lleva un par de filas por unidad de origen —lo alimentado y su ley— en vez
    de una sola fila agregada. Sin ese detalle no se sabe de donde viene lo que
    entra, y la ley promedio de alimentacion no se puede recalcular.
    """
    origenes = [u for u in unidades if u.entrega_a == unidad.nombre]
    fila = fila_de_seccion(hoja, fila, "Alimentacion recibida", anos)
    for origen in origenes:
        numericas.append(fila)
        fila = fila_de_entrada(
            hoja, fila, f"Concentrado alimentado desde {origen.nombre}", "t", anos, True
        )
        # La ley que interesa es la de los metales que el complejo refina, no la
        # de todos los que produce el origen: Pisco es una fundicion de estano y
        # el concentrado de cobre de una unidad polimetalica se vende aparte.
        for metal in unidad.metales:
            numericas.append(fila)
            fila = fila_de_entrada(
                hoja, fila, f"Ley de {metal} del concentrado de {origen.nombre}", "%", anos
            )

    fila = fila_de_seccion(hoja, fila, "Complejo", anos)
    for concepto, medida, corroborable in COMPLEJO:
        numericas.append(fila)
        fila = fila_de_entrada(hoja, fila, concepto, medida, anos, corroborable)
    for metal in unidad.metales:
        for plantilla, medida, corroborable in COMPLEJO_POR_METAL:
            numericas.append(fila)
            fila = fila_de_entrada(
                hoja, fila, plantilla.format(metal=metal), medida, anos, corroborable
            )

    # El libro distingue la recuperacion de SR + B2 de la de NZ + SRP: no hay
    # una sola para todo el complejo. Que criterio agrupa esta consultado a
    # Finanzas; mientras tanto se pide una por unidad de origen, que es mas
    # general y reproduce el libro dando el mismo valor a las de un grupo.
    fila = fila_de_seccion(hoja, fila, "Recuperacion por origen", anos)
    for origen in origenes:
        for metal in unidad.metales:
            numericas.append(fila)
            fila = fila_de_entrada(
                hoja, fila, f"Recuperacion de {metal} de {origen.nombre}", "%", anos
            )
    return fila, numericas


def _corriente(
    hoja: Worksheet,
    unidad: Unidad,
    corriente: tuple[str, str | None, str, bool],
    fila: int,
    numericas: list[int],
    anos: int,
) -> tuple[int, list[int]]:
    """Un tonelaje y, debajo, la ley de cada metal que lleva."""
    concepto, _, sufijo_de_ley, corroborable = corriente
    numericas.append(fila)
    fila = fila_de_entrada(hoja, fila, concepto, "t", anos, corroborable)
    for metal in unidad.metales:
        numericas.append(fila)
        fila = fila_de_entrada(
            hoja, fila, f"Ley de {metal} {sufijo_de_ley}", "%", anos, corroborable
        )
    return fila, numericas


def _nombre_de_hoja(nombre: str) -> str:
    """Excel limita el nombre de una hoja a 31 caracteres y prohibe varios."""
    limpio = "".join(c for c in nombre if c not in r"[]:*?/\'")
    return limpio[:31] or "Unidad"


def _descripcion(unidad: Unidad) -> str:
    partes = [f"origen: {unidad.origen}", f"metales: {', '.join(unidad.metales)}"]
    if unidad.etapas:
        partes.append(f"etapas: {', '.join(unidad.etapas)}")
    if unidad.entrega_a:
        partes.append(f"entrega a: {unidad.entrega_a}")
    return " · ".join(partes)


def hoja_opex(libro: Workbook, unidades: list[Unidad], primer_ano: int, anos: int) -> None:
    hoja = libro.create_sheet("Opex")
    encabezar(hoja, "Costo operativo por unidad productiva, en US$", primer_ano, anos)
    fila, numericas = 5, []
    for unidad in unidades:
        fila = fila_de_seccion(hoja, fila, f"{unidad.nombre} ({unidad.tipo})", anos)
        for concepto, etapa in OPEX_POR_UNIDAD:
            if not unidad.aplica(etapa):
                continue
            numericas.append(fila)
            fila = fila_de_entrada(hoja, fila, concepto, "US$", anos)
        # El acuerdo 6 de la minuta del 27/08/2026 exige que la lista sea
        # extensible: estas filas quedan para conceptos propios del proyecto.
        for _ in range(3):
            numericas.append(fila)
            fila = fila_de_entrada(hoja, fila, "", "US$", anos)
        fila += 1
    validacion_numerica(hoja, numericas, anos)


def hoja_capex(libro: Workbook, unidades: list[Unidad], primer_ano: int, anos: int) -> None:
    hoja = libro.create_sheet("Capex")
    encabezar(hoja, "Capital por unidad, etapa y naturaleza contable, en US$", primer_ano, anos)
    fila, numericas = 5, []
    for unidad in unidades:
        fila = fila_de_seccion(hoja, fila, f"{unidad.nombre} ({unidad.tipo})", anos)
        for etapa in CAPEX_ETAPAS:
            fila = fila_de_seccion(hoja, fila, f"  {etapa}", anos)
            for naturaleza in CAPEX_NATURALEZAS:
                numericas.append(fila)
                fila = fila_de_entrada(hoja, fila, f"    {naturaleza}", "US$", anos)
        fila += 1
    validacion_numerica(hoja, numericas, anos)


def hoja_precios(libro: Workbook, unidades: list[Unidad], primer_ano: int, anos: int) -> None:
    hoja = libro.create_sheet("Precios")
    encabezar(hoja, "Precios y terminos comerciales por metal", primer_ano, anos)
    metales = sorted({metal for unidad in unidades for metal in unidad.metales})
    fila, numericas = 5, []
    for metal in metales:
        fila = fila_de_seccion(hoja, fila, metal, anos)
        for escenario in ESCENARIOS_DE_PRECIO:
            numericas.append(fila)
            fila = fila_de_entrada(hoja, fila, f"Precio, escenario {escenario}", "US$/t", anos)
        for concepto, medida in (
            # El libro vende el metal por dos caminos: refinado al precio mas
            # un premio, y en concentrado al precio por el factor pagable.
            ("Premio del metal refinado", "US$/t"),
            ("Factor de metal pagable", "fraccion"),
            ("Terminos comerciales", "US$/t"),
            ("Gasto de ventas", "US$/t"),
            ("Fletes", "US$/t"),
        ):
            numericas.append(fila)
            fila = fila_de_entrada(hoja, fila, concepto, medida, anos)
        fila += 1
    validacion_numerica(hoja, numericas, anos)


def hoja_instrucciones(libro: Workbook, unidades: list[Unidad]) -> None:
    hoja = libro.create_sheet("Leeme")
    lineas = [
        ("Plantilla canonica de inputs", TITULO),
        ("", None),
        ("Las celdas con color son las que se llenan. El resto es estructura.", None),
        ("Una celda vacia significa que el concepto no aplica ese ano.", None),
        ("", None),
        ("Las celdas verdes son corroborables", CABECERA),
        ("Se llenan igual que las crema. La diferencia es que el sistema las", None),
        ("recalcula a partir del resto de la cadena y avisa si el dato cargado no", None),
        ("cuadra. El dato cargado es el que se usa: el recalculo lo audita, no lo", None),
        ("sustituye.", None),
        ("", None),
        ("Una pestana por proyecto", CABECERA),
        ("Cada unidad tiene su hoja, con los sub-bloques Mina y Planta: de donde", None),
        ("sale el mineral y por que proceso pasa. Cada corriente de tonelaje lleva", None),
        ("debajo la ley de cada metal que transporta.", None),
        ("", None),
        ("La estructura la determinan las unidades declaradas al generar la", None),
        ("plantilla. Si el caso necesita otra unidad, se vuelve a generar; no se", None),
        ("agregan filas a mano, porque el motor lee la estructura, no el formato.", None),
        ("", None),
        ("Unidades de este caso:", CABECERA),
    ]
    for unidad in unidades:
        detalle = [
            unidad.tipo,
            f"origen: {unidad.origen}",
            f"metales: {', '.join(unidad.metales)}",
        ]
        if unidad.etapas:
            detalle.append(f"etapas: {', '.join(unidad.etapas)}")
        if unidad.entrega_a:
            detalle.append(f"entrega a: {unidad.entrega_a}")
        lineas.append((f"  {unidad.nombre} — {' — '.join(detalle)}", None))
    lineas += [
        ("", None),
        ("Confidencialidad", CABECERA),
        ("Llenada con datos de un caso real, esta plantilla es informacion", None),
        ("confidencial de MINSUR y no vuelve al repositorio de codigo.", None),
        ("", None),
        ("Referencia", CABECERA),
        ("docs/modelo-economico/catalogo-inputs.md define cada concepto.", None),
        ("docs/modelo-economico/modelo-estandar.md explica las entidades.", None),
    ]
    for i, (texto, fuente) in enumerate(lineas, start=1):
        celda = hoja.cell(row=i, column=1, value=texto)
        if fuente is not None:
            celda.font = fuente
    hoja.column_dimensions["A"].width = 80


def _encadenar(unidades: list[Unidad]) -> list[Unidad]:
    """Dirige el concentrado de cada mina a la fundicion, si el caso declara una.

    No se pide como campo aparte porque el caso admite una sola fundicion, y en
    el libro las cinco minas entregan a Pisco. Un proyecto que venda su
    concentrado directo simplemente no declara fundicion.
    """
    fundiciones = [u for u in unidades if u.es_fundicion]
    if len(fundiciones) > 1:
        raise ValueError(
            "El caso declara mas de una fundicion. El tope de capacidad se aplica sobre el "
            "concentrado del complejo y no sabria a cual acotar."
        )
    if not fundiciones:
        return unidades
    destino = fundiciones[0].nombre
    return [u if u.es_fundicion else replace(u, entrega_a=destino) for u in unidades]


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--salida", type=Path, required=True, help="Ruta del .xlsx a escribir")
    p.add_argument(
        "--unidad",
        action="append",
        required=True,
        metavar="NOMBRE:TIPO:METALES[:ETAPAS][:ORIGEN]",
        help=f"Unidad productiva. TIPO en {{{', '.join(TIPOS_DE_UNIDAD)}}}",
    )
    p.add_argument("--primer-ano", type=int, required=True)
    p.add_argument("--anos", type=int, default=36)
    p.add_argument(
        "--bloque",
        choices=("produccion", "completo"),
        default="produccion",
        help="Bloque a emitir. 'completo' mantiene el libro unico anterior.",
    )
    args = p.parse_args()

    try:
        unidades = [Unidad.desde_texto(u) for u in args.unidad]
    except ValueError as error:
        print(f"Declaracion de unidad invalida: {error}")
        return 1

    if not 1 <= args.anos <= 60:
        print("El horizonte debe estar entre 1 y 60 anos.")
        return 1

    try:
        unidades = _encadenar(unidades)
    except ValueError as error:
        print(f"Declaracion de unidad invalida: {error}")
        return 1

    libro = Workbook()
    hoja_caso(libro, unidades, args.primer_ano, args.anos)
    for unidad in unidades:
        hoja_produccion_de_unidad(libro, unidad, unidades, args.primer_ano, args.anos)
    if args.bloque == "completo":
        # Opex, capex y precios siguen en hojas unicas con las unidades
        # apiladas. Les toca su propia plantilla mas adelante; hasta entonces
        # el libro completo es el que lee la ingesta de punta a punta.
        hoja_opex(libro, unidades, args.primer_ano, args.anos)
        hoja_capex(libro, unidades, args.primer_ano, args.anos)
        hoja_precios(libro, unidades, args.primer_ano, args.anos)
    hoja_instrucciones(libro, unidades)

    args.salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(args.salida)

    print(f"Plantilla escrita en {args.salida}")
    print(f"  {len(unidades)} unidad(es), {args.anos} anos desde {args.primer_ano}")
    print(f"  hojas: {', '.join(libro.sheetnames)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
