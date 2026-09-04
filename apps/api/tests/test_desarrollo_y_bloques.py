"""Pruebas de la carga local y de los bloques intermedios.

Cubren el recorrido que hace la interfaz de modelamiento de escenario: crear el
caso, subir la plantilla, calcular y recorrer la cadena hoja por hoja.

Y cubren sobre todo el limite: la carga por la API existe solo en local, y los
datos maestros de desarrollo tambien. Si esa frontera se borrara sin que nadie
lo notara, un entorno del cliente acabaria calculando con parametros que MINSUR
no ha confirmado, que es exactamente lo que `R-32` impide.
"""

from __future__ import annotations

import importlib.util
import sys
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from minsur_api import maestros_desarrollo
from minsur_api.bloques import BASE_DEL_CASH_COST, _grupo_de_la_refineria
from minsur_api.dependencias import datos_maestros, repositorio
from minsur_api.main import crear_app
from minsur_api.repositorio import RepositorioEnMemoria
from minsur_engine.refineria import AporteALaRefineria, BloqueDeLaRefineria
from minsur_ingest.capex import CON_DATO_DE_CAPEX
from minsur_ingest.opex import CON_DATO_DE_GASTOS

CABECERAS = {"Authorization": "Bearer token-de-desarrollo"}
RAIZ = Path(__file__).resolve().parents[3]
RUTA_DE_CARGA = "/api/desarrollo/casos/CASO-XX-2026-000000/insumos"

# La cadena de una mina, en el orden en que la plantilla pide las filas. El
# primer ejercicio invierte y no produce.
# La cadena de una mina de estano, con las magnitudes del oficio: las leyes en
# por ciento tal como se teclean en la hoja, y cada eslabon cuadrando con el
# anterior. Las cifras son inventadas y no salen de MINSUR, pero **el orden de
# magnitud importa**: con una ley de cabeza de centesimas de punto la pantalla
# redondeaba todas las leyes a `0.0%` y no habia forma de ver si la cadena
# estaba bien.
#
#     finas       = tratado x ley del tratado x recuperacion
#                 = 420 000 x 4,2 % x 90 %      = 15 876 t
#     concentrado = finas / ley del concentrado
#                 = 15 876 / 60 %               = 26 460 t
CADENA = [
    1_200_000.0,  # Mineral extraido
    1.80,  # Ley Sn de cabeza
    1_200_000.0,  # Mineral tratado en preconcentracion
    1.80,  # Ley de Sn de entrada
    420_000.0,  # Mineral preconcentrado
    4.20,  # Ley Sn del preconcentrado
    0.0,  # Mineral directo
    0.0,  # Ley de Sn del directo
    420_000.0,  # Mineral tratado total
    4.20,  # Ley Sn del tratado total
    420_000.0,  # Mineral tratado total para cash cost
    4.20,  # Ley Sn del cash cost
    15_876.0,  # Toneladas finas
    60.0,  # Ley Sn del concentrado
    90.0,  # Recuperacion Sn
    26_460.0,  # Produccion Concentrado
    0.0,  # Concentrado producido Cu
    0.0,  # Ley Cu
    0.0,  # Ley Ag
]


def _generador() -> object:
    """Importa el script que emite la plantilla, que no es un paquete."""
    ruta = RAIZ / "scripts" / "generar_plantilla_inputs.py"
    especificacion = importlib.util.spec_from_file_location("generador", ruta)
    assert especificacion is not None and especificacion.loader is not None
    modulo = importlib.util.module_from_spec(especificacion)
    sys.modules["generador"] = modulo
    especificacion.loader.exec_module(modulo)
    return modulo


def _escribir(hoja: object, etiqueta: str, valores: list[float]) -> None:
    for fila in hoja.iter_rows(min_row=5, max_col=1):  # type: ignore[attr-defined]
        celda = fila[0]
        if celda.value and str(celda.value).strip() == etiqueta.strip():
            for i, valor in enumerate(valores):
                hoja.cell(row=celda.row, column=3 + i, value=valor)  # type: ignore[attr-defined]
            return
    raise AssertionError(f"la plantilla no tiene la fila {etiqueta!r}")


def _llenar_cadena(hoja: object, anos: int = 3) -> None:
    """Llena la pestaña por posición, que es como se lee."""
    fila = 5
    for valor in CADENA:
        while not hoja.cell(row=fila, column=2).value:  # type: ignore[attr-defined]
            fila += 1
        for i in range(anos):
            hoja.cell(row=fila, column=3 + i, value=0.0 if i == 0 else valor)  # type: ignore[attr-defined]
        fila += 1


@pytest.fixture
def plantilla(tmp_path: Path) -> Path:
    """Un libro de una mina y una refinería, con la cadena llena."""
    generador = _generador()
    ruta = tmp_path / "caso.xlsx"
    unidades = generador._encadenar(  # type: ignore[attr-defined]
        [
            generador.Unidad.desde_texto("Mina Alfa:mina:Sn"),  # type: ignore[attr-defined]
            generador.Unidad.desde_texto("Refineria:refineria:Sn"),  # type: ignore[attr-defined]
        ]
    )
    libro = generador.Workbook()  # type: ignore[attr-defined]
    generador.hoja_caso(libro, unidades, 2027, 3)  # type: ignore[attr-defined]
    generador.hoja_produccion_de_unidad(libro, "Mina Alfa", 2027, 3)  # type: ignore[attr-defined]
    generador.hoja_precios(libro, unidades, 2027, 3)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    libro["Caso"]["B3"] = "Escenario de prueba"
    libro["Caso"]["B9"] = "CP-2026-09"
    _llenar_cadena(libro["Mina Alfa"])
    _escribir(libro["Precios"], "Precio, escenario Base", [10_000.0] * 3)
    _escribir(libro["Precios"], "Factor de metal pagable", [0.9] * 3)
    libro.save(ruta)
    return ruta


@pytest.fixture
def plantilla_opex(tmp_path: Path) -> Path:
    """El libro de opex, con una pestaña por unidad y dos conceptos llenos."""
    generador = _generador()
    ruta = tmp_path / "opex.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    for nombre in ("Mina Alfa", "Refineria"):
        generador.hoja_opex_de_unidad(libro, nombre, 2027, 3)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    for nombre in ("Mina Alfa", "Refineria"):
        _escribir(libro[nombre], "Mina", [0.0, 90_000.0, 90_000.0])
        _escribir(libro[nombre], "Gastos administrativos", [0.0, 12_000.0, 12_000.0])
    libro.save(ruta)
    return ruta


@pytest.fixture
def plantilla_capex(tmp_path: Path) -> Path:
    """El libro de capital, con dos naturalezas llenas en cada unidad."""
    generador = _generador()
    ruta = tmp_path / "capex.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    for nombre in ("Mina Alfa", "Refineria"):
        generador.hoja_capex_de_unidad(libro, nombre, 2027, 3)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    for nombre in ("Mina Alfa", "Refineria"):
        _escribir(libro[nombre], "No depreciable", [0.0, 0.0, 400.0])
        _escribir(libro[nombre], "Equipos de cómputo", [30.0, 0.0, 0.0])
        _escribir(libro[nombre], "Maquinaria, equipos y vehículos", [9_000.0, 1_500.0, 0.0])
    libro.save(ruta)
    return ruta


@pytest.fixture
def repo() -> RepositorioEnMemoria:
    return RepositorioEnMemoria()


@pytest.fixture
def cliente(repo: RepositorioEnMemoria) -> Iterator[TestClient]:
    app = crear_app()
    app.dependency_overrides[repositorio] = lambda: repo
    with TestClient(app) as instancia:
        yield instancia
    app.dependency_overrides.clear()


def _crear_caso(cliente: TestClient) -> str:
    respuesta = cliente.post(
        "/api/casos",
        json={
            "abreviatura": "PRB-1",
            "nombre": "Escenario de prueba",
            "tipo": "con-proyecto",
            "descripcion": "Prueba",
        },
        headers=CABECERAS,
    )
    assert respuesta.status_code == 201, respuesta.text
    return str(respuesta.json()["id_caso"])


def _subir(
    cliente: TestClient,
    id_caso: str,
    ruta: Path,
    opex: Path | None = None,
    capex: Path | None = None,
) -> dict[str, object]:
    abiertos = [ruta.open("rb")]
    archivos = {"caso": (ruta.name, abiertos[0], "application/vnd.ms-excel")}
    for clave, libro in (("opex", opex), ("capex", capex)):
        if libro is None:
            continue
        abiertos.append(libro.open("rb"))
        archivos[clave] = (libro.name, abiertos[-1], "application/vnd.ms-excel")
    try:
        respuesta = cliente.post(
            f"/api/desarrollo/casos/{id_caso}/insumos", files=archivos, headers=CABECERAS
        )
    finally:
        for archivo in abiertos:
            archivo.close()
    assert respuesta.status_code == 200, respuesta.text
    resultado: dict[str, object] = respuesta.json()
    return resultado


class TestLaFronteraDeLoLocal:
    def test_en_local_los_datos_maestros_son_los_de_desarrollo(self) -> None:
        maestros = datos_maestros()
        assert maestros is not None
        assert maestros.version == maestros_desarrollo.VERSION == "DEV-0"

    def test_fuera_de_local_no_hay_datos_maestros(self, monkeypatch: pytest.MonkeyPatch) -> None:
        # Es la garantia que sostiene `R-32`: los parametros los confirma MINSUR
        # y ningun entorno del cliente puede calcular sin ellos.
        monkeypatch.setattr("minsur_api.dependencias.config", lambda: _ConfigRemota())
        assert datos_maestros() is None

    def test_fuera_de_local_la_carga_por_la_api_no_existe(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Se comprueba por comportamiento y no por introspeccion: la ruta esta
        # fuera del esquema a proposito, y `app.routes` conserva los routers sin
        # desplegar en esta version de FastAPI. Un 404 es la ausencia; un 422
        # seria la ruta existiendo y rechazando el cuerpo.
        monkeypatch.setattr("minsur_api.main.config", lambda: _ConfigRemota())
        with TestClient(crear_app()) as remoto:
            respuesta = remoto.post(RUTA_DE_CARGA, headers=CABECERAS)
        assert respuesta.status_code == 404

    def test_en_local_la_carga_esta_montada_y_fuera_del_contrato(self, cliente: TestClient) -> None:
        assert cliente.post(RUTA_DE_CARGA, headers=CABECERAS).status_code == 422

        # De este esquema salen los tipos del frontend y la definicion que se
        # importa a API Management. Una ruta que solo existe en la maquina del
        # desarrollador no pertenece a ese contrato.
        rutas = crear_app().openapi()["paths"]
        assert not [ruta for ruta in rutas if "desarrollo" in ruta]
        assert "/api/casos" in rutas, "el resto de la API sigue publicada"


class _ConfigRemota:
    """Lo mínimo que las dependencias consultan de la configuración."""

    entorno = "qa"
    es_local = False


class TestCargaDeInsumos:
    def test_la_plantilla_llena_deja_el_caso_calculable(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        id_caso = _crear_caso(cliente)
        resultado = _subir(cliente, id_caso, plantilla)

        assert resultado["valida"] is True, resultado["hallazgos"]
        assert resultado["incidencias"] == []
        assert int(str(resultado["filas_leidas"])) > 0
        assert len(str(resultado["huella"])) == 64

        respuesta = cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        assert respuesta.status_code == 200, respuesta.text
        assert respuesta.json()["terna"]["datos_maestros"] == "DEV-0"

    def test_la_revision_de_inputs_sube_con_cada_carga(
        self, cliente: TestClient, repo: RepositorioEnMemoria, plantilla: Path
    ) -> None:
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla)
        primera = repo.obtener(id_caso).revision_inputs
        _subir(cliente, id_caso, plantilla)
        assert repo.obtener(id_caso).revision_inputs == primera + 1

    def test_un_archivo_ilegible_devuelve_la_incidencia_con_su_hoja(
        self, cliente: TestClient, tmp_path: Path
    ) -> None:
        id_caso = _crear_caso(cliente)
        roto = tmp_path / "roto.xlsx"
        roto.write_bytes(b"esto no es un libro de Excel")
        with roto.open("rb") as archivo:
            respuesta = cliente.post(
                f"/api/desarrollo/casos/{id_caso}/insumos",
                files={"caso": (roto.name, archivo, "application/vnd.ms-excel")},
                headers=CABECERAS,
            )
        assert respuesta.status_code == 200, respuesta.text
        cuerpo = respuesta.json()
        assert cuerpo["valida"] is False
        assert cuerpo["incidencias"], "una lectura fallida tiene que decir por que"
        assert set(cuerpo["incidencias"][0]) >= {"hoja", "celda", "mensaje"}

    def test_un_caso_inexistente_no_admite_carga(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        with plantilla.open("rb") as archivo:
            respuesta = cliente.post(
                "/api/desarrollo/casos/CASO-XX-2026-000000/insumos",
                files={"caso": (plantilla.name, archivo, "application/vnd.ms-excel")},
                headers=CABECERAS,
            )
        assert respuesta.status_code == 404


class TestBloquesIntermedios:
    def _bloques(
        self,
        cliente: TestClient,
        plantilla: Path,
        opex: Path | None = None,
        capex: Path | None = None,
    ) -> dict[str, Any]:
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla, opex, capex)
        cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        respuesta = cliente.get(f"/api/casos/{id_caso}/corrida/bloques", headers=CABECERAS)
        assert respuesta.status_code == 200, respuesta.text
        cuerpo: dict[str, Any] = respuesta.json()
        return cuerpo

    def test_los_bloques_van_en_el_orden_del_libro(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        cuerpo = self._bloques(cliente, plantilla)

        # `Depreciacion` va antes que `Ventas` porque asi esta en el libro,
        # aunque la cadena de calculo no lo exija. Es el orden aprendido.
        assert [b["clave"] for b in cuerpo["bloques"]] == [
            "produccion",
            "opex",
            "capex",
            "depreciacion",
            "ventas",
            "otros",
            "impuestos",
            "flujo",
        ]

    def test_el_complejo_no_es_una_hoja_aparte(self, cliente: TestClient, plantilla: Path) -> None:
        # En el libro Pisco es un bloque dentro de `InputsProd`, no una hoja. Lo
        # fue una version anterior de la pantalla y no debe volver a serlo.
        cuerpo = self._bloques(cliente, plantilla)
        assert "refineria" not in [b["clave"] for b in cuerpo["bloques"]]

        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        assert [g["titulo"] for g in produccion["grupos"]][-2:] == ["Refineria", "Venta Sn Spot"]

    def test_la_hoja_cierra_con_la_venta_spot_y_su_check(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        cuerpo = self._bloques(cliente, plantilla)
        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        spot = produccion["grupos"][-1]
        etiquetas = [s["etiqueta"] for s in spot["secciones"][0]["series"]]
        assert etiquetas == [
            "Concentrado Excedente",
            "Ley Promedio de Alimentación",
            "Producción Sn Refinado",
            "Check",
        ]
        # El `Check` se cumple por construccion, asi que la fila lo dice: quien
        # la lea no debe tomarla por una verificacion activa.
        assert spot["secciones"][0]["series"][-1]["nota"]

    def test_produccion_usa_el_vocabulario_del_libro(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # Las etiquetas y las medidas salen del catalogo de la ingesta, que es
        # con el que se emite la plantilla y con el que se lee. El usuario ve en
        # pantalla la fila que lleno, con su nombre y su unidad.
        cuerpo = self._bloques(cliente, plantilla)
        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        mina = next(g for g in produccion["grupos"] if g["titulo"] == "Mina Alfa")

        assert [s["titulo"] for s in mina["secciones"]][:2] == ["Mina", "Planta"]
        primera = mina["secciones"][0]["series"][0]
        assert primera["etiqueta"] == "Mineral extraído"
        assert primera["medida"] == "t"
        assert primera["concepto"] == "mineral_extraido"
        assert primera["origen"] == "dato"

        leyes = [s for s in mina["secciones"][1]["series"] if s["etiqueta"].startswith("Ley")]
        assert leyes and all(s["medida"] in {"%", "oz/t"} for s in leyes)

    def test_opex_sigue_el_orden_de_la_hoja(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # Primero el cash cost de cada unidad con su total, despues lo que el
        # libro calcula debajo, y los gastos al final. Los gastos no van dentro
        # del bloque de la unidad: no son cash cost y cada fila va a un sitio
        # distinto del flujo.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        titulos = [g["titulo"] for g in opex["grupos"]]

        assert titulos[0].startswith("Cash Cost - ")
        assert "Producción" in titulos
        assert titulos[-1] == "Total Gastos"
        assert any(t.startswith("Gastos - ") for t in titulos)
        assert titulos.index("Producción") > max(
            i for i, t in enumerate(titulos) if t.startswith("Cash Cost - ")
        )

    def test_la_estructura_base_del_cash_cost(self) -> None:
        # Los diez conceptos que MINSUR quiere ver en todos los bloques, en su
        # orden. Se toman del catalogo por posicion, asi que esta prueba es la
        # que avisa si alguien lo reordena: sin ella el cambio pasaria en
        # silencio y la hoja mostraria otra cosa.
        assert BASE_DEL_CASH_COST == (
            "Exploraciones",
            "Geología",
            "Mina",
            "Planta Preconcentración",
            "Planta Concentradora",
            "Mantenimiento",
            "Energía",
            "Apoyo",
            "Estudios y optimizaciones",
            "Relavera",
        )

    def test_la_base_aparece_aunque_este_en_cero(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # La plantilla de la prueba solo llena `Mina`. Las otras nueve de la
        # base tienen que salir igual, con su guion.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        primero = next(g for g in opex["grupos"] if g["titulo"].startswith("Cash Cost - "))
        etiquetas = [s["etiqueta"] for s in primero["secciones"][0]["series"]]

        assert etiquetas[: len(BASE_DEL_CASH_COST)] == list(BASE_DEL_CASH_COST)
        assert etiquetas[-1].startswith("Total ")

    def test_lo_que_no_es_de_la_base_se_oculta_si_esta_vacio(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # Es lo que devuelve la lista de conceptos de cada unidad: el bloque de
        # una mina no lleva las filas de la refineria.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        primero = next(g for g in opex["grupos"] if g["titulo"].startswith("Cash Cost - "))
        etiquetas = {s["etiqueta"] for s in primero["secciones"][0]["series"]}

        assert "Fundición" not in etiquetas
        assert "Línea de transmisión" not in etiquetas

    def test_el_reparto_de_la_refineria_se_informa_y_no_entra_al_flujo(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # Regla 079. El supuesto `Costo de Fundicion` se pedia en la plantilla y
        # la ingesta lo descartaba; ahora llega al motor. El bloque lo reparte
        # por origen, pero no se suma al cash cost: que origenes van por el
        # bloque directo lo tiene que declarar el caso, y sumarlo antes cobraria
        # dos veces la tonelada de uno que ya esta dentro.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        supuestos = [g for g in opex["grupos"] if g["titulo"].startswith("Supuestos ")]

        # La plantilla de la prueba no declara la tarifa, asi que el bloque no
        # se emite: sin tarifa no hay reparto que mostrar.
        assert supuestos == []

    def test_la_unidad_con_costo_directo_no_paga_tarifa(self) -> None:
        # La bandera la declara cada unidad en su pestana de supuestos, y con
        # ella el bloque directo cubre a unas y la tarifa a las otras. Que
        # unidades son no puede estar escrito en el codigo: un proyecto nuevo
        # entra por la tarifa hasta que alguien decida lo contrario.
        from minsur_engine.caso import ProduccionDeUnidad, UnidadProductiva

        sin_declarar = UnidadProductiva(
            nombre="Proyecto X", tipo="mina", produccion=ProduccionDeUnidad(mineral_tratado=(1.0,))
        )
        assert sin_declarar.costo_directo_en_la_refineria is False

    def test_el_total_del_cash_cost_va_en_su_propia_fila(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # En el libro cierra los bloques de cash cost y no entra en el de
        # produccion: es una fila suelta entre los dos.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        produccion = next(g for g in opex["grupos"] if g["titulo"] == "Producción")

        assert "Total Cash Cost" not in [
            s["etiqueta"] for seccion in produccion["secciones"] for s in seccion["series"]
        ]
        suelta = next(
            g
            for g in opex["grupos"]
            if any(
                s["etiqueta"] == "Total Cash Cost"
                for seccion in g["secciones"]
                for s in seccion["series"]
            )
        )
        assert opex["grupos"].index(suelta) < opex["grupos"].index(produccion)

    def test_el_costo_por_tonelada_tratada_se_calcula_una_sola_vez(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # El modelo no lo hace por unidad: lo calcula sobre la unidad cuyo cash
        # cost se analiza por tonelada tratada, y la plataforma lo ancla a la
        # primera del caso que trate mineral.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        unitario = next(
            g for g in opex["grupos"] if g["titulo"] == "Cash cost por tonelada tratada"
        )

        assert len(unitario["secciones"]) == 1
        series = unitario["secciones"][0]["series"]
        assert all(s["medida"] == "$/tt" for s in series)
        totales = [s["etiqueta"] for s in series if s["etiqueta"].startswith("Total ")]
        assert len(totales) == 1, "un solo total, no uno por unidad"

    def test_los_gastos_muestran_su_estructura_completa(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # A diferencia del cash cost, la lista de gastos es la misma para todas
        # las unidades: ocultar lo vacio no devolveria la lista de nadie y solo
        # movería cada concepto de linea segun el proyecto.
        #
        # Y es la misma que la de la tabla de cierre, las dos filas derivadas
        # incluidas. Con listas distintas, quien comparaba una unidad contra el
        # total contaba lineas distintas.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        primero = next(g for g in opex["grupos"] if g["titulo"].startswith("Gastos - "))
        cierre = next(g for g in opex["grupos"] if g["titulo"] == "Total Gastos")
        series = primero["secciones"][0]["series"]
        etiquetas = [s["etiqueta"] for s in series]

        assert etiquetas == [s["etiqueta"] for s in cierre["secciones"][0]["series"]]
        del_catalogo = [f.etiqueta for f in CON_DATO_DE_GASTOS]
        assert [e for e in etiquetas if e in del_catalogo] == del_catalogo
        assert etiquetas.index("Planilla") == etiquetas.index("Servidumbres y usufructos") + 1
        assert etiquetas[-1] == "Gestión Social Deducible"
        assert any(all(v == 0 for v in s["valores"]) for s in series)

    def test_la_hoja_cierra_con_el_total_de_gastos(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        # El modelo cierra con los gastos consolidados, y ese bloque lleva las
        # dos filas que el motor deriva y la plantilla no pide: la planilla y la
        # gestion social deducible. Reglas 026 y 027.
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")

        assert opex["grupos"][-1]["titulo"] == "Total Gastos"
        etiquetas = [s["etiqueta"] for s in opex["grupos"][-1]["secciones"][0]["series"]]
        assert "Año con operación" not in etiquetas
        assert "Planilla" in etiquetas
        assert etiquetas[-1] == "Gestión Social Deducible"

    def _capex(self, cliente: TestClient, plantilla: Path, capex: Path) -> dict[str, Any]:
        cuerpo = self._bloques(cliente, plantilla, capex=capex)
        bloque: dict[str, Any] = next(b for b in cuerpo["bloques"] if b["clave"] == "capex")
        return bloque

    def test_capex_sigue_el_orden_de_la_hoja(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El libro abre con la etapa consolidada, sigue con la clasificacion
        # contable de cada unidad -lo unico que la hoja carga- y cierra con el
        # cruce de las dos clasificaciones y su total.
        titulos = [g["titulo"] for g in self._capex(cliente, plantilla, plantilla_capex)["grupos"]]

        assert titulos[0] == "Detalle Capex"
        assert titulos[-1] == "Capex por Naturaleza TOTAL"
        assert [t for t in titulos if t.startswith("Clasificación ")] == [
            "Clasificación Mina Alfa",
            "Clasificación Refineria",
        ]
        assert titulos.index("Capex por Naturaleza") > titulos.index("Clasificación Refineria")

    def test_la_clasificacion_contable_va_completa(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Las cinco naturalezas en todas las unidades, con su guion en la fila
        # vacia: la lista es la misma para todas y lo que importa es que cada
        # concepto caiga siempre en la misma linea.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        grupo = next(g for g in bloque["grupos"] if g["titulo"].startswith("Clasificación "))
        series = grupo["secciones"][0]["series"]

        assert [s["etiqueta"] for s in series[:-1]] == [f.etiqueta for f in CON_DATO_DE_CAPEX]
        assert series[-1]["etiqueta"] == "Total"
        assert series[-1]["total"] is True
        assert all(s["origen"] == "dato" for s in series[:-1])
        assert any(not any(s["valores"]) for s in series[:-1]), "alguna naturaleza va en cero"

    def test_el_cruce_lleva_las_cinco_naturalezas_por_etapa(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El libro muestra cuatro filas porque el computo comparte codigo con la
        # maquinaria. Aqui van cinco: lo que llega sumado no se separa.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        grupo = next(g for g in bloque["grupos"] if g["titulo"] == "Capex por Naturaleza")

        assert [s["titulo"] for s in grupo["secciones"]] == [
            "Por Naturaleza - Inicial",
            "Por Naturaleza - Sostenimiento",
            "Por Naturaleza - Cierre",
        ]
        for seccion in grupo["secciones"]:
            etiquetas = [s["etiqueta"] for s in seccion["series"]]
            assert etiquetas == [f.etiqueta for f in CON_DATO_DE_CAPEX] + ["Total"]
        computo = next(
            s for s in grupo["secciones"][2]["series"] if s["concepto"] == "equipos_de_computo"
        )
        assert computo["nota"], "la fila dice que el libro la lleva sumada"

    def test_la_hoja_cierra_con_su_cuadre_en_cero(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Con una sola clasificacion cargada se cumple por construccion. Se
        # muestra igual, como indicador, y si dejara de cerrar seria senal de
        # que el capital que se pinta no es el que entro al calculo.
        #
        # El libro trae ademas `Tipo vs Detalle`, que aqui no se emite: restaba
        # el total de la suma de los totales por unidad, y las dos cifras salen
        # de las mismas celdas.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        series = [
            serie
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        ]
        assert [s["etiqueta"] for s in series].count("Tipo vs Detalle") == 0

        cuadre = next(s for s in series if s["etiqueta"] == "check")
        assert cuadre["valores"] == pytest.approx([0.0] * len(cuadre["valores"]))
        assert cuadre["nota"], "la fila declara que se cumple por construccion"

    def test_las_naturalezas_llevan_el_codigo_contable_del_libro(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # La columna A del libro. No es decorativa: es la clave de union de los
        # `SUMIF` que construyen todo lo derivado, y que dos filas compartan
        # `MAQ` es lo que explica que el libro fusione el computo con la
        # maquinaria justo donde la plataforma los separa.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        clasificacion = next(
            g for g in bloque["grupos"] if g["titulo"].startswith("Clasificación ")
        )
        series = clasificacion["secciones"][0]["series"]

        assert [s["codigo"] for s in series] == ["NOD", "MAQ", "MAQ", "INS", "EDI", ""]

        cruce = next(g for g in bloque["grupos"] if g["titulo"] == "Capex por Naturaleza")
        assert [s["codigo"] for s in cruce["secciones"][0]["series"][:5]] == [
            "NOD",
            "MAQ",
            "MAQ",
            "INS",
            "EDI",
        ]
        # El consolidado final no lo lleva, y el libro tampoco: ahi no se agrupa
        # nada, se suman las tres etapas que los `SUMIF` ya construyeron.
        total = next(g for g in bloque["grupos"] if g["titulo"] == "Capex por Naturaleza TOTAL")
        assert all(not s["codigo"] for s in total["secciones"][0]["series"])

    def test_el_capital_se_muestra_en_miles(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El libro lleva esta hoja en `$k` y la ingesta convierte a dolares al
        # leer. La pantalla deshace esa conversion por la unidad de medida, de
        # modo que rotularla `US$` mostraba la hoja mil veces mas grande.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        medidas = {
            serie["medida"]
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        }
        assert medidas == {"$k"}

    def test_la_etapa_del_resumen_es_la_del_cruce(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Las dos zonas dicen lo mismo desde sitios distintos: el resumen lo toma
        # del capital que entro al calculo y el cruce lo deriva con la regla de
        # etapa del motor. Si difieren, el capital que se pinta no es el que se
        # deprecio, y ninguna de las dos filas de cuadre lo notaria: las dos
        # comparan totales, y el total coincide aunque el reparto no.
        bloque = self._capex(cliente, plantilla, plantilla_capex)
        resumen = {
            s["etiqueta"]: s["valores"] for s in bloque["grupos"][0]["secciones"][0]["series"]
        }
        cruce = next(g for g in bloque["grupos"] if g["titulo"] == "Capex por Naturaleza")

        for etiqueta, seccion in zip(
            ("Capex Inicial", "Sostenimiento", "Cierre Mina"), cruce["secciones"], strict=True
        ):
            assert seccion["series"][-1]["etiqueta"] == "Total"
            assert resumen[etiqueta] == pytest.approx(seccion["series"][-1]["valores"]), etiqueta

    def test_sin_unidad_de_proyecto_no_hay_bloque_por_proyecto(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El libro desdobla el cruce para el proyecto que evalua. Unidad de
        # proyecto es la que declara umbral de capital inicial, y este caso no
        # lo declara: el consolidado va solo, sin bloque suelto que lo duplique.
        titulos = [g["titulo"] for g in self._capex(cliente, plantilla, plantilla_capex)["grupos"]]
        assert [t for t in titulos if t.startswith("Capex por Naturaleza - ")] == []

    def _depreciacion(self, cliente: TestClient, plantilla: Path, capex: Path) -> dict[str, Any]:
        cuerpo = self._bloques(cliente, plantilla, capex=capex)
        bloque: dict[str, Any] = next(b for b in cuerpo["bloques"] if b["clave"] == "depreciacion")
        return bloque

    def test_depreciacion_sigue_el_orden_de_la_hoja(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El libro lleva las dos vias seguidas y no intercaladas: primero toda la
        # tributaria con su consolidado, despues toda la financiera, y al final
        # el capital que queda sin depreciar.
        titulos = [
            g["titulo"] for g in self._depreciacion(cliente, plantilla, plantilla_capex)["grupos"]
        ]

        assert titulos[0].startswith("Depreciación Tributaria - ")
        assert titulos[-1] == "Capital y depreciación acumulada"
        tributario = titulos.index("Total Depreciación Tributaria")
        financiero = titulos.index("Total Depreciación Financiera")
        assert tributario < financiero
        assert all(
            tributario < titulos.index(t) < financiero
            for t in titulos
            if t.startswith("Depreciación Financiera - ")
        )

    def test_el_resumen_tributario_lleva_las_lineas_del_libro(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # El resumen del libro va **por etapa** y no por naturaleza contable: dice
        # cuanto viene del capital inicial y cuanto del de sostenimiento.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        grupo = next(
            g for g in bloque["grupos"] if g["titulo"].startswith("Depreciación Tributaria - ")
        )
        resumen = grupo["secciones"][0]

        assert resumen["titulo"] is None
        assert [s["etiqueta"] for s in resumen["series"]] == [
            "Proyección SAP",
            "Depreciación - Capex Inicial",
            "Depreciación - Sostenimiento",
            "Escudo Fiscal - Cierre de Mina",
            "Depreciación Estudios",
            "Total Depreciación",
            "Periodo con Producción",
        ]
        assert resumen["series"][5]["total"] is True

    def test_los_triangulos_de_cosechas_vienen_plegados(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Son tres cuartas partes de la hoja: una fila por ano de inversion. El
        # resumen no se pliega, porque es lo que se lee primero.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        grupo = next(
            g for g in bloque["grupos"] if g["titulo"].startswith("Depreciación Tributaria - ")
        )
        plegables = [s for s in grupo["secciones"] if s["plegable"]]

        assert grupo["secciones"][0]["plegable"] is False
        assert plegables, "el bloque abre al menos un triangulo"
        for seccion in plegables:
            assert " · " in (seccion["titulo"] or ""), "el rotulo es etapa y componente"
            assert seccion["series"][-1]["etiqueta"] == "Total"
            assert seccion["series"][-1]["total"] is True

    def test_la_via_financiera_lleva_su_tabla_de_agotamiento(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Sin ella la cuota financiera es un numero sin derivacion, y es justo
        # donde el contraste contra el modelo no cierra.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        grupo = next(
            g for g in bloque["grupos"] if g["titulo"].startswith("Depreciación Financiera - ")
        )
        agotamiento = next(s for s in grupo["secciones"] if s["titulo"] == "Agotamiento")

        assert [s["etiqueta"] for s in agotamiento["series"]] == [
            "Reservas",
            "Mineral Extraído",
            "Conversión Recursos",
            "Reservas Finales",
            "Tasa Depreciación Capex",
            "Capex (sin maquinarias)",
            "Saldo Inicial",
            "Depreciación Capex",
            "Saldo Final",
        ]

    def test_un_saldo_no_lleva_total_del_horizonte(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # La unidad de medida no basta para decidirlo: un saldo va en las mismas
        # unidades que el flujo que lo mueve. Sumar los saldos de apertura de
        # treinta y seis ejercicios da una cifra que no significa nada, y en la
        # pantalla se veia tres veces el capital del caso.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        series = {
            serie["etiqueta"]: serie
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        }

        for etiqueta in (
            "Reservas",
            "Reservas Finales",
            "Saldo Inicial",
            "Saldo Final",
            "Acum. Capex no depreciado",
        ):
            assert series[etiqueta]["acumulado"] is None, etiqueta
        # Los flujos de la misma tabla sí lo llevan.
        assert series["Depreciación Capex"]["acumulado"] is not None
        assert series["Mineral Extraído"]["acumulado"] is not None

    def test_el_consolidado_financiero_cuadra_con_las_unidades(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # `Nuevo Capex` no es una linea del desglose sino las tres etapas juntas,
        # y sin ella el consolidado salia en cero teniendo cada unidad su cuota.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        por_unidad = [
            serie["valores"]
            for grupo in bloque["grupos"]
            if grupo["titulo"].startswith("Depreciación Financiera - ")
            for serie in grupo["secciones"][0]["series"]
            if serie["etiqueta"] == "Total Depreciación Financiera"
        ]
        consolidado = next(
            g for g in bloque["grupos"] if g["titulo"] == "Total Depreciación Financiera"
        )
        total = next(
            s
            for s in consolidado["secciones"][0]["series"]
            if s["etiqueta"] == "Total Depreciación Financiera"
        )

        assert por_unidad, "hay al menos una unidad con depreciacion financiera"
        sumado = [sum(v[i] for v in por_unidad) for i in range(len(total["valores"]))]
        assert total["valores"] == pytest.approx(sumado)

    def test_la_depreciacion_se_muestra_en_miles(
        self, cliente: TestClient, plantilla: Path, plantilla_capex: Path
    ) -> None:
        # Como el libro y como las otras dos hojas de dinero. Antes iba en `US$` y
        # se veia mil veces mas grande que el capital que la origina.
        bloque = self._depreciacion(cliente, plantilla, plantilla_capex)
        medidas = {
            serie["medida"]
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        }
        # La bandera del periodo no lleva unidad, y la tabla de agotamiento
        # ademas tonelaje y una tasa.
        assert medidas <= {"$k", "t", "%", ""}
        assert "$k" in medidas

    def test_ventas_sigue_el_orden_de_la_hoja(self, cliente: TestClient, plantilla: Path) -> None:
        # El libro abre con las cuatro lineas de venta y sus precios, sigue con
        # la liquidacion del concentrado y cierra con lo que vende cada unidad.
        cuerpo = self._bloques(cliente, plantilla)
        ventas = next(b for b in cuerpo["bloques"] if b["clave"] == "ventas")
        grupo = ventas["grupos"][0]

        assert grupo["titulo"] == "Ventas Netas"
        assert [s["titulo"] for s in grupo["secciones"]] == [
            "Ventas Sn refinado",
            "Ventas Sn Spot",
            None,
            "Ventas concentrado Cu",
            None,
        ]
        assert [s["etiqueta"] for s in grupo["secciones"][0]["series"]] == [
            "Volumen Sn",
            "Precio total",
            "Precio Spot Sn",
            "Premio Sn",
            "Venta Sn Refinado",
        ]

    def test_la_ranura_reservada_no_se_emite(self, cliente: TestClient, plantilla: Path) -> None:
        # La fila 26 del libro, rotulada `xxx`, es la segunda de sus tres ranuras
        # reservadas: sin formula y en cero siempre. Es la regla `042`.
        cuerpo = self._bloques(cliente, plantilla)
        ventas = next(b for b in cuerpo["bloques"] if b["clave"] == "ventas")
        etiquetas = [
            serie["etiqueta"]
            for grupo in ventas["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        ]

        assert "xxx" not in etiquetas
        assert "Venta Total" in etiquetas
        assert "Ajustes finales" in etiquetas

    def test_la_venta_se_muestra_en_miles(self, cliente: TestClient, plantilla: Path) -> None:
        # Como el libro y como las otras hojas de dinero. Antes iba en `US$`.
        cuerpo = self._bloques(cliente, plantilla)
        ventas = next(b for b in cuerpo["bloques"] if b["clave"] == "ventas")
        medidas = {
            serie["medida"]
            for grupo in ventas["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        }

        assert "$k" in medidas
        assert "US$" not in medidas

    def test_la_ley_se_totaliza_ponderada_por_su_tonelaje(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # El modelo da total tambien en las leyes, y no es una suma: sumar
        # treinta y seis porcentajes no es nada. Es el promedio ponderado por el
        # tonelaje de su fila, que es como el propio libro consolida la ley de
        # dos corrientes. La relacion entre una ley y su tonelaje es posicional.
        cuerpo = self._bloques(cliente, plantilla)
        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        mina = next(g for g in produccion["grupos"] if g["titulo"] == "Mina Alfa")
        series = [s for seccion in mina["secciones"] for s in seccion["series"]]

        extraido = next(s for s in series if s["etiqueta"] == "Mineral extraído")
        ley = next(s for s in series if s["etiqueta"] == "Ley Sn")

        assert extraido["acumulado"] == pytest.approx(sum(extraido["valores"]))
        # Con la ley constante en los ejercicios con dato, el ponderado es esa
        # misma ley y no su suma.
        assert ley["acumulado"] == pytest.approx(max(ley["valores"]))
        assert ley["acumulado"] < sum(ley["valores"])

    def test_las_leyes_de_la_refineria_tambien_se_totalizan(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # Toda ley del bloque lleva su total, ponderado por el tonelaje de su
        # fila: la del concentrado de un origen por lo que ese origen entrega, y
        # la recuperacion tambien. Una media sin pesar daria el mismo valor a
        # una unidad que aporta el ochenta por ciento y a otra que aporta el dos.
        cuerpo = self._bloques(cliente, plantilla)
        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        # Por su contenido y no por su rotulo: el nombre de la unidad lo pone el
        # caso, y aqui la refineria se llama `Refineria`.
        series = [
            serie
            for grupo in produccion["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
            if serie["etiqueta"].startswith(("Concentrado ", "Ley ", "Recuperación Sn "))
        ]

        leyes = [s for s in series if s["medida"] == "%"]
        assert leyes, "el bloque de la refineria tiene leyes"
        assert all(s["acumulado"] is not None for s in leyes), [
            s["etiqueta"] for s in leyes if s["acumulado"] is None
        ]

    def test_la_ley_de_la_refineria_se_pondera_por_lo_entregado(self) -> None:
        # Cuando la refineria satura, lo alimentado es menor que lo entregado, y
        # la ley consolidada del libro se divide entre lo **entregado**: es el
        # denominador de su formula, `(l1*c1 + l2*c2) / (c1 + c2)` sobre las
        # filas por origen, que no llevan tope.
        #
        # Pesando por lo alimentado, un ejercicio en que la refineria no recibe
        # nada dejaria la ley del horizonte en cero aunque cada ejercicio tuviera
        # la suya, que es justo lo que se veia en pantalla.
        # Dos ejercicios de cien toneladas entregadas y leyes distintas, con el
        # primero saturado del todo: sin el, los dos pesos darian lo mismo y la
        # prueba no distinguiria nada.
        #
        #     por lo entregado   (0,60 x 100 + 0,40 x 100) / 200 = 0,50
        #     por lo alimentado  (0,60 x   0 + 0,40 x 100) / 100 = 0,40
        aporte = AporteALaRefineria(
            unidad="Mina Alfa",
            concentrado=(100.0, 100.0),
            ley=(0.60, 0.40),
            recuperacion=(0.90, 0.90),
            a_spot=(100.0, 0.0),
            refinado=(0.0, 36.0),
            refinado_sin_restriccion=(54.0, 36.0),
        )
        bloque = BloqueDeLaRefineria(
            aportes=(aporte,),
            concentrado_entregado=(100.0, 100.0),
            concentrado_alimentado=(0.0, 100.0),
            ley_de_alimentacion=(0.60, 0.40),
            toneladas_alimentadas=(0.0, 100.0),
            refinado=(0.0, 36.0),
            refinado_sin_restriccion=(54.0, 36.0),
            concentrado_excedente=(100.0, 0.0),
            ley_del_excedente=(0.60, 0.0),
            refinado_del_excedente=(60.0, 0.0),
            check=(0.0, 0.0),
        )
        series = {
            s.etiqueta: s for s in _grupo_de_la_refineria(bloque, "Pisco").secciones[0].series
        }

        assert series["Ley de Sn en Concentrado"].acumulado == pytest.approx(0.50)
        assert series["Ley Promedio de Alimentación"].acumulado == pytest.approx(0.50)

    def test_un_ratio_no_lleva_total(
        self, cliente: TestClient, plantilla: Path, plantilla_opex: Path
    ) -> None:
        cuerpo = self._bloques(cliente, plantilla, plantilla_opex)
        opex = next(b for b in cuerpo["bloques"] if b["clave"] == "opex")
        por_fina = next(g for g in opex["grupos"] if g["titulo"] == "Cash cost por tonelada fina")

        assert all(s["acumulado"] is None for s in por_fina["secciones"][0]["series"])

    def test_cada_serie_tiene_un_valor_por_ano(self, cliente: TestClient, plantilla: Path) -> None:
        cuerpo = self._bloques(cliente, plantilla)
        anios = len(cuerpo["anios"])
        assert anios == 3
        for bloque in cuerpo["bloques"]:
            for grupo in bloque["grupos"]:
                for seccion in grupo["secciones"]:
                    for serie in seccion["series"]:
                        assert len(serie["valores"]) == anios, (
                            bloque["clave"],
                            serie["etiqueta"],
                        )

    def test_la_produccion_trae_su_recalculo_para_contrastar(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # Es la alerta de control de calidad: sin el valor recalculado al lado,
        # una diferencia no dice que esperaba el sistema. Las ocho filas que lo
        # llevan son las que el catalogo marca como calculadas.
        cuerpo = self._bloques(cliente, plantilla)
        produccion = next(b for b in cuerpo["bloques"] if b["clave"] == "produccion")
        mina = next(g for g in produccion["grupos"] if g["titulo"] == "Mina Alfa")
        series = [s for seccion in mina["secciones"] for s in seccion["series"]]

        assert [s["etiqueta"] for s in series if s["recalculada"] is not None]
        assert all(s["origen"] == "dato" for s in series)

    def test_los_campos_con_dato_los_decide_el_motor(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla)
        cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        cuerpo = cliente.get(f"/api/casos/{id_caso}/corrida/bloques", headers=CABECERAS).json()

        assert "Mina Alfa" in cuerpo["campos_con_dato"]
        assert "mineral_extraido" in cuerpo["campos_con_dato"]["Mina Alfa"]

    def test_un_caso_sin_corrida_no_tiene_bloques(self, cliente: TestClient) -> None:
        id_caso = _crear_caso(cliente)
        respuesta = cliente.get(f"/api/casos/{id_caso}/corrida/bloques", headers=CABECERAS)
        assert respuesta.status_code == 409


class TestLaHojaOtros:
    """La pestana `Otros`, con la forma del libro.

    Es la hoja que arma la bolsa de egresos, el IGV y el capital de trabajo, y
    la que reparte la regalia entre gasto y tributo. Hasta el 04/09/2026 emitia
    doce series planas en dos secciones y en dolares.
    """

    def _otros(self, cliente: TestClient, plantilla: Path) -> dict[str, Any]:
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla)
        cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        cuerpo = cliente.get(f"/api/casos/{id_caso}/corrida/bloques", headers=CABECERAS).json()
        bloque: dict[str, Any] = next(b for b in cuerpo["bloques"] if b["clave"] == "otros")
        return bloque

    def _series(self, bloque: dict[str, Any]) -> list[dict[str, Any]]:
        return [
            serie
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        ]

    def test_sigue_el_orden_de_las_bandas_del_libro(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        bloque = self._otros(cliente, plantilla)

        assert [g["titulo"] for g in bloque["grupos"]] == [
            "Cash Cost",
            "Gasto de Ventas",
            "Fletes",
            "Otros Gastos/ Ingresos",
            "Impuestos",
            "Compras",
            "Δ WK",
            "Cuentas por Cobrar",
            "Cuentas por Pagar",
            "",
            "Tipo de Compra/Venta",
            "",
            "Otros Flujo de Caja",
        ]

    def test_la_bolsa_lleva_los_once_conceptos_del_libro(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # El libro no lleva la gestion social en este bloque y el motor la
        # sumaba: es la regla `081`. Doce filas de concepto serian esa regresion.
        bloque = self._otros(cliente, plantilla)
        compras = next(g for g in bloque["grupos"] if g["titulo"] == "Compras")
        etiquetas = [s["etiqueta"] for s in compras["secciones"][0]["series"]]

        assert etiquetas == [
            "Opex",
            "Gastos Administrativos",
            "Fletes",
            "Gasto de Ventas",
            "Donaciones",
            "Otros Egresos",
            "Servidumbre",
            "Estudios",
            "Planilla",
            "Capex",
            "Exploraciones",
            "Bolsa Egresos",
        ]
        assert "Gestión Social" not in etiquetas

    def test_la_hoja_se_muestra_en_miles(self, cliente: TestClient, plantilla: Path) -> None:
        # El libro la lleva en `$k` y la ingesta convierte a dolares al leer. La
        # pantalla deshace esa conversion por la unidad de medida, de modo que
        # rotularla `US$` mostraba la hoja mil veces mas grande.
        bloque = self._otros(cliente, plantilla)
        medidas = {serie["medida"] for serie in self._series(bloque)}

        assert medidas == {"$k", "%", "dias", ""}

    def test_los_saldos_no_llevan_acumulado(self, cliente: TestClient, plantilla: Path) -> None:
        # Sumar los saldos de apertura de treinta y seis ejercicios da una cifra
        # que no significa nada. Las cinco filas de saldo lo declaran.
        bloque = self._otros(cliente, plantilla)
        saldos = {"CXC", "CxP", "Credito acumulado", "CxC otros", "CxP otros"}
        por_etiqueta = {s["etiqueta"]: s for s in self._series(bloque)}

        assert saldos <= set(por_etiqueta)
        for etiqueta in saldos:
            assert por_etiqueta[etiqueta]["acumulado"] is None, etiqueta

    def test_la_regalia_aparece_en_los_dos_bloques_con_su_rotulo(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # `Otros!33` y `!37` son la misma cifra repartida por el signo de la
        # utilidad operativa, y los dos rotulos del libro difieren solo en la
        # tilde. Cada uno lleva la nota que dice donde esta el otro.
        bloque = self._otros(cliente, plantilla)
        gastos = next(g for g in bloque["grupos"] if g["titulo"] == "Otros Gastos/ Ingresos")
        tributos = next(g for g in bloque["grupos"] if g["titulo"] == "Impuestos")

        con_tilde = next(s for s in gastos["secciones"][0]["series"] if s["etiqueta"] == "Regalías")
        sin_tilde = next(
            s for s in tributos["secciones"][0]["series"] if s["etiqueta"] == "Regalias"
        )
        assert con_tilde["nota"]
        assert sin_tilde["nota"]

    def test_las_dos_filas_de_presentacion_no_se_emiten(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # La `40`, rotulada `xxx`, es la tercera ranura reservada del libro. La
        # `58` repite bajo el nombre `IGV Ventas Locales` la variacion que la
        # misma banda trae nueve filas mas abajo.
        etiquetas = [s["etiqueta"] for s in self._series(self._otros(cliente, plantilla))]

        assert "xxx" not in etiquetas
        assert "IGV Ventas Locales" not in etiquetas
        assert "Variación IGV Flujo Caja" in etiquetas

    def test_el_igv_se_ve_entero_y_llega_al_flujo_en_cero(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # Las dos mitades de la regla `014`. La variacion es un `@property` del
        # motor y por eso no salia cuando el bloque se emitia por reflexion.
        bloque = self._otros(cliente, plantilla)
        igv = next(g for g in bloque["grupos"] if g["titulo"] == "Δ WK")
        por_etiqueta = {s["etiqueta"]: s for s in igv["secciones"][0]["series"]}

        assert "Credito/Pago" in por_etiqueta
        variacion = por_etiqueta["Variación IGV Flujo Caja"]
        assert all(v == 0.0 for v in variacion["valores"])
        assert variacion["nota"]

    def test_cada_banda_cierra_en_la_suma_de_sus_filas(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # El total lo trae el motor y la pantalla no lo recalcula: si las filas
        # que se pintan no lo suman, lo que se ve y lo que alimenta el flujo se
        # han separado sin que nada lo acuse.
        bloque = self._otros(cliente, plantilla)
        cerradas = 0
        for grupo in bloque["grupos"]:
            for seccion in grupo["secciones"]:
                series = seccion["series"]
                if not series[-1]["total"]:
                    continue
                cerradas += 1
                for i, obtenido in enumerate(series[-1]["valores"]):
                    esperado = sum(s["valores"][i] for s in series[:-1])
                    assert obtenido == pytest.approx(esperado), f"{grupo['titulo']}, ano {i}"
        # Las diez bandas que cierran: seis de egreso y tributo, el IGV, las dos
        # cuentas comerciales y las no comerciales.
        assert cerradas == 10


class TestLaEscalaDeLaHojaImpuestos:
    def test_el_dinero_va_en_miles_y_las_tasas_en_por_ciento(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # `Otros` e `Impuestos` publican la misma regalia. Con esta hoja en
        # dolares, la cifra se veia mil veces mas grande en una que en otra.
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla)
        cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        cuerpo = cliente.get(f"/api/casos/{id_caso}/corrida/bloques", headers=CABECERAS).json()
        bloque = next(b for b in cuerpo["bloques"] if b["clave"] == "impuestos")
        por_concepto = {
            serie["concepto"]: serie
            for grupo in bloque["grupos"]
            for seccion in grupo["secciones"]
            for serie in seccion["series"]
        }

        assert por_concepto["regalia_mayor"]["medida"] == "$k"
        assert por_concepto["tasa_impuesto_renta"]["medida"] == "%"
        assert por_concepto["margen_operativo"]["medida"] == "%"
        # La perdida arrastrada abre y cierra en un saldo, y un saldo no se suma.
        assert por_concepto["saldo_inicial"]["acumulado"] is None
        assert por_concepto["saldo_final"]["acumulado"] is None


class TestLaTirNoFingeUnCero:
    def test_un_caso_sin_desembolso_declara_que_no_tiene_tir(
        self, cliente: TestClient, plantilla: Path
    ) -> None:
        # La plantilla no declara capital, asi que el flujo no abre con
        # desembolso y no hay tasa que describa una rentabilidad. El libro
        # escribe un guion; el contrato escribe nulo. Lo decide el ADR 0011.
        id_caso = _crear_caso(cliente)
        _subir(cliente, id_caso, plantilla)
        respuesta = cliente.post(f"/api/casos/{id_caso}/evaluar", headers=CABECERAS)
        assert respuesta.status_code == 200, respuesta.text
        assert respuesta.json()["indicadores"]["tir"] is None
