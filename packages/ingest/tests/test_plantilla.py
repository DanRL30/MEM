"""Pruebas de la lectura de plantillas.

La prueba que sostiene el paquete es la de ida y vuelta: se genera la plantilla
con el script que la emite, se llena, se lee y se comprueba que sale lo que se
escribió. Generador y lector son las dos mitades del mismo contrato, y esta
prueba falla en cuanto una se mueve sin la otra.

Las cifras son las del libro de producción de MINSUR: 1 000 t extraídas al 2 %,
de las que 600 pasan por preconcentración al 1,2 % y salen 300 t al 2,4 %. El
directo son 400 t al 3,2 % por balance; el tratado total, 700 t al 2,857 %; y con
90 % de recuperación y un concentrado al 40 %, 18 t finas y 45 t de concentrado.
"""

from __future__ import annotations

import importlib.util
import math
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from minsur_engine.caso import DatosMaestros
from minsur_engine.corrida import calcular
from minsur_engine.depreciacion import TasasDeDepreciacion
from minsur_engine.impuestos import EscalaProgresiva, Tramo
from minsur_engine.parametros import ParametrosCorporativos
from minsur_ingest.incidencias import ErrorDePlantilla
from minsur_ingest.plantilla import (
    Lectura,
    leer_comite_de_precios,
    leer_o_fallar,
    leer_plantilla,
    leer_produccion,
    leer_supuestos,
)
from minsur_ingest.produccion import (
    FILAS_CON_DATO,
    UNIDADES_PROVISIONALES,
    ErrorDeAsociacion,
    asociar_por_orden,
)
from minsur_ingest.sinonimos import canonizar
from minsur_ingest.supuestos import aplicar as aplicar_supuestos

RAIZ = Path(__file__).resolve().parents[3]

MAESTROS = DatosMaestros(
    parametros=ParametrosCorporativos(
        version_datos_maestros="CP-PRUEBA",
        tasa_descuento=0.10,
        participacion_trabajadores=0.08,
        impuesto_renta=0.295,
        regalia_minima=0.01,
        osinergmin=0.0,
        oefa=0.0,
        fondo_jubilacion_minera=0.005,
    ),
    tasas_tributarias=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    tasas_financieras=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    escala_regalia=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.01),)),
    escala_iem=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.0),)),
)

# La cadena, en el orden en que la estructura estandar la pide y en las unidades
# de la plantilla: las leyes en por ciento y las onzas por tonelada tal cual. Se
# llena por posicion porque el libro repite la etiqueta `Ley Sn` cinco veces, que
# es justo la razon de leer la plantilla por secuencia y no por nombre.
CADENA = [
    1_000.0,  # Mineral extraido
    2.0,  # Ley Sn de cabeza
    600.0,  # Mineral Tratado en Pre Concentracion
    1.2,  # Ley de Sn (entrada)
    300.0,  # Mineral Pre-Concentrado a Concentradora
    2.4,  # Ley Sn del preconcentrado
    400.0,  # Mineral Directo a Planta Concentradora
    3.2,  # Ley de Sn del directo
    700.0,  # Mineral Tratado Total en Concentradora
    2.857142857142857,  # Ley Sn del tratado total
    1_000.0,  # Mineral Tratado Total (Cash Cost)
    2.0,  # Ley Sn del cash cost
    18.0,  # Toneladas finas
    40.0,  # Ley Sn Concentrado
    90.0,  # Recuperacion Sn
    45.0,  # Produccion Concentrado
    0.0,  # Concentrado Producido Cu
    0.0,  # Ley Cu
    0.0,  # Ley Ag
]


def _generador() -> object:
    """Importa el script que emite la plantilla, que no es un paquete."""
    ruta = RAIZ / "scripts" / "generar_plantilla_inputs.py"
    especificacion = importlib.util.spec_from_file_location("generador", ruta)
    assert especificacion is not None and especificacion.loader is not None
    modulo = importlib.util.module_from_spec(especificacion)
    sys.modules["generador"] = modulo
    especificacion.loader.exec_module(modulo)
    return modulo


def _escribir(hoja: object, etiqueta: str, valores: list[float]) -> int:
    """Llena la fila de un concepto y devuelve su número."""
    for fila in hoja.iter_rows(min_row=5, max_col=1):  # type: ignore[attr-defined]
        celda = fila[0]
        if celda.value and str(celda.value).strip() == etiqueta.strip():
            for i, valor in enumerate(valores):
                hoja.cell(row=celda.row, column=3 + i, value=valor)  # type: ignore[attr-defined]
            return int(celda.row)
    raise AssertionError(f"la plantilla no tiene la fila {etiqueta!r}")


def _llenar_cadena(hoja: object, anos: int = 3) -> None:
    """Llena la pestaña por posición, que es como se lee.

    El primer ejercicio invierte y no produce; los siguientes repiten la cadena.
    """
    fila = 5
    for valor in CADENA:
        while not hoja.cell(row=fila, column=2).value:  # type: ignore[attr-defined]
            fila += 1
        for i in range(anos):
            hoja.cell(row=fila, column=3 + i, value=0.0 if i == 0 else valor)  # type: ignore[attr-defined]
        fila += 1


@pytest.fixture
def plantilla_llena(tmp_path: Path) -> Path:
    """Un libro completo de una mina y una refinería, con la cadena llena."""
    generador = _generador()
    ruta = tmp_path / "caso.xlsx"
    unidades = generador._encadenar(  # type: ignore[attr-defined]
        [
            generador.Unidad.desde_texto("Mina Alfa:mina:Sn"),  # type: ignore[attr-defined]
            generador.Unidad.desde_texto("Refineria:refineria:Sn"),  # type: ignore[attr-defined]
        ]
    )
    libro = generador.Workbook()  # type: ignore[attr-defined]
    generador.hoja_caso(libro, unidades, 2027, 3)  # type: ignore[attr-defined]
    generador.hoja_produccion_de_unidad(libro, "Mina Alfa", 2027, 3)  # type: ignore[attr-defined]
    generador.hoja_precios(libro, unidades, 2027, 3)  # type: ignore[attr-defined]
    generador.hoja_instrucciones(libro, unidades)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    caso = libro["Caso"]
    caso["B3"] = "Caso de prueba"
    caso["B9"] = "CP-2026-09"
    _llenar_cadena(libro["Mina Alfa"])

    precios = libro["Precios"]
    _escribir(precios, "Precio, escenario Base", [10_000.0] * 3)
    _escribir(precios, "Factor de metal pagable", [0.9] * 3)
    libro.save(ruta)
    return ruta


class TestIdaYVuelta:
    def test_la_plantilla_llena_produce_un_caso(self, plantilla_llena: Path) -> None:
        lectura = leer_plantilla(plantilla_llena)
        assert lectura.valida, [str(i) for i in lectura.incidencias]
        caso = lectura.caso
        assert caso is not None
        assert caso.nombre == "Caso de prueba"
        assert caso.horizonte.primer_ano == 2027
        assert [u.nombre for u in caso.unidades] == ["Mina Alfa", "Refineria"]

    def test_la_cadena_completa_llega_al_motor(self, plantilla_llena: Path) -> None:
        # Hasta el 01/09/2026 el lector consumia cinco series por unidad y
        # descartaba el resto. Esta prueba fija que la cadena entera llega.
        p = leer_o_fallar(plantilla_llena).unidades[0].produccion
        assert p.mineral_extraido == (0.0, 1_000.0, 1_000.0)
        assert p.ley_de_cabeza == (0.0, 0.02, 0.02)
        assert p.tratado_en_preconcentracion == (0.0, 600.0, 600.0)
        assert p.preconcentrado == (0.0, 300.0, 300.0)
        assert p.directo == (0.0, 400.0, 400.0)
        assert p.ley_del_directo == (0.0, 0.032, 0.032)
        assert p.tratado_total == (0.0, 700.0, 700.0)
        assert p.mineral_tratado == (0.0, 1_000.0, 1_000.0)
        assert p.toneladas_finas == (0.0, 18.0, 18.0)
        assert p.recuperacion == (0.0, 0.90, 0.90)
        assert p.concentrado_producido == (0.0, 45.0, 45.0)

    def test_la_refineria_no_necesita_pestana(self, plantilla_llena: Path) -> None:
        # Sus filas son resultado de lo que producen las minas: pedirlas como
        # dato invitaria a que contradijeran a su origen.
        caso = leer_o_fallar(plantilla_llena)
        assert caso.refineria is not None
        assert not any(caso.refineria.produccion.concentrado_producido)

    def test_el_origen_y_el_destino_sobreviven_la_ida_y_vuelta(self, plantilla_llena: Path) -> None:
        alfa = leer_o_fallar(plantilla_llena).unidades[0]
        assert alfa.origen == "yacimiento"
        assert alfa.entrega_a == "Refineria"

    def test_el_caso_leido_calcula_y_no_reporta_discrepancias(self, plantilla_llena: Path) -> None:
        # La prueba de que la frontera funciona y de que la cadena cargada es
        # coherente: lo que sale de la ingesta entra en el motor sin retoques y
        # el recalculo no encuentra nada que advertir.
        corrida = calcular(leer_o_fallar(plantilla_llena), MAESTROS)
        assert corrida.discrepancias == ()
        assert corrida.mineral_tratado_por_unidad["Mina Alfa"] == (0.0, 1_000.0, 1_000.0)

    def test_un_dato_alterado_se_advierte_y_el_calculo_sigue(self, plantilla_llena: Path) -> None:
        # Es alarma y control de calidad, no correccion: el motor avisa y llega
        # igual hasta el NPV para que se vea el efecto.
        libro = load_workbook(plantilla_llena)
        _escribir(libro["Mina Alfa"], "Producción Concentrado", [0.0, 60.0, 45.0])
        libro.save(plantilla_llena)

        corrida = calcular(leer_o_fallar(plantilla_llena), MAESTROS)
        assert [d.concepto for d in corrida.discrepancias] == ["Produccion Concentrado"]
        assert corrida.discrepancias[0].recalculado == pytest.approx(45.0)


class TestDatosComunesDelCaso:
    """La hoja `Caso` recoge lo que no es de ninguna unidad."""

    def _escribir_comun(self, ruta: Path, etiqueta: str, valor: str | float) -> None:
        libro = load_workbook(ruta)
        hoja = libro["Caso"]
        for fila in hoja.iter_rows(min_row=1, max_col=3):
            if fila[0].value and canonizar(str(fila[0].value)) == canonizar(etiqueta):
                hoja.cell(row=fila[0].row, column=3, value=valor)
                break
        else:
            raise AssertionError(f"no existe la fila {etiqueta!r} en la hoja Caso")
        libro.save(ruta)

    def test_el_interruptor_de_cuentas_comerciales_llega_al_caso(
        self, plantilla_llena: Path
    ) -> None:
        # `Control!$G$21` del libro. Hasta ahora el campo existia, el motor lo
        # implementaba y ninguna plantilla podia escribirlo.
        self._escribir_comun(plantilla_llena, "Cuentas comerciales en el capital de trabajo", "no")
        lectura = leer_plantilla(plantilla_llena)
        assert lectura.caso is not None, [str(i) for i in lectura.incidencias]
        assert lectura.caso.datos_comunes.cuentas_de_capital_trabajo_activas is False

    def test_el_interruptor_sin_llenar_deja_las_cuentas_encendidas(
        self, plantilla_llena: Path
    ) -> None:
        lectura = leer_plantilla(plantilla_llena)
        assert lectura.caso is not None
        assert lectura.caso.datos_comunes.cuentas_de_capital_trabajo_activas is True

    def test_una_bandera_que_no_es_si_ni_no_se_reporta(self, plantilla_llena: Path) -> None:
        self._escribir_comun(
            plantilla_llena, "Cuentas comerciales en el capital de trabajo", "quizas"
        )
        lectura = leer_plantilla(plantilla_llena)
        assert any("si o no" in i.mensaje for i in lectura.incidencias)

    def test_un_dato_comun_con_valor_que_nadie_consume_se_reporta(
        self, plantilla_llena: Path
    ) -> None:
        # El fallo sin sintoma: se llena, se lee sin errores y no interviene.
        # Hoy le pasa a los costos hundidos, declarados en cuatro documentos y
        # ausentes del codigo.
        self._escribir_comun(plantilla_llena, "Costos hundidos excluidos del flujo", 1_000.0)
        lectura = leer_plantilla(plantilla_llena)
        assert any("todavia no lo consume" in i.mensaje for i in lectura.incidencias)

    def test_un_dato_comun_sin_consumidor_y_vacio_no_molesta(self, plantilla_llena: Path) -> None:
        # Una celda vacia es legitima: la plantilla es la misma para todos.
        lectura = leer_plantilla(plantilla_llena)
        assert not [i for i in lectura.incidencias if "todavia no lo consume" in i.mensaje]

    def test_los_dias_de_rotacion_ya_no_salen_de_esta_hoja(self, plantilla_llena: Path) -> None:
        # El libro los lleva en dos filas distintas y aqui habia una sola celda
        # que llenaba las dos con el mismo numero. Viven en supuestos.
        libro = load_workbook(plantilla_llena)
        etiquetas = {
            canonizar(str(fila[0].value))
            for fila in libro["Caso"].iter_rows(min_row=1, max_col=1)
            if fila[0].value
        }
        assert canonizar("Dias de working capital") not in etiquetas
        lectura = leer_plantilla(plantilla_llena)
        assert lectura.caso is not None
        assert lectura.caso.datos_comunes.dias_por_cobrar == ()


class TestConversionDeEscalas:
    def test_la_ley_de_plata_no_se_lee_como_porcentaje(self, plantilla_llena: Path) -> None:
        # Viene en onzas troy por tonelada. Tratarla como porcentaje la
        # dividiria entre cien sin avisar.
        libro = load_workbook(plantilla_llena)
        _escribir(libro["Mina Alfa"], "Ley Ag", [0.0, 10.5, 10.5])
        libro.save(plantilla_llena)
        assert leer_o_fallar(plantilla_llena).unidades[0].produccion.ley_ag == (0.0, 10.5, 10.5)


class TestIncidencias:
    def test_una_celda_con_texto_se_reporta_con_su_ubicacion(self, plantilla_llena: Path) -> None:
        libro = load_workbook(plantilla_llena)
        fila = _escribir(libro["Mina Alfa"], "Mineral extraído", [0.0] * 3)
        libro["Mina Alfa"].cell(row=fila, column=4, value="mil toneladas")
        libro.save(plantilla_llena)

        lectura = leer_plantilla(plantilla_llena)
        assert not lectura.valida
        incidencia = lectura.incidencias[0]
        assert incidencia.hoja == "Mina Alfa"
        assert incidencia.celda == f"D{fila}"
        assert "se esperaba un numero" in incidencia.mensaje

    def test_una_estructura_alterada_se_reporta_y_no_se_adivina(
        self, plantilla_llena: Path
    ) -> None:
        # Si la secuencia se rompe, seguir leyendo asignaria cada serie al
        # concepto de al lado y el caso saldria plausible y equivocado.
        libro = load_workbook(plantilla_llena)
        libro["Mina Alfa"].cell(row=6, column=1, value="Mineral flotado")
        libro.save(plantilla_llena)

        lectura = leer_plantilla(plantilla_llena)
        assert not lectura.valida
        assert any("estructura fija" in i.mensaje for i in lectura.incidencias)

    def test_una_unidad_declarada_sin_pestana_se_reporta(self, plantilla_llena: Path) -> None:
        libro = load_workbook(plantilla_llena)
        del libro["Mina Alfa"]
        libro.save(plantilla_llena)

        lectura = leer_plantilla(plantilla_llena)
        assert not lectura.valida
        assert any("no tiene pestana" in i.mensaje for i in lectura.incidencias)

    def test_un_libro_que_no_es_la_plantilla_se_rechaza(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        ruta = tmp_path / "otra_cosa.xlsx"
        Workbook().save(ruta)
        lectura = leer_plantilla(ruta)
        assert not lectura.valida
        assert "no es la canonica" in lectura.incidencias[0].mensaje

    def test_un_archivo_que_no_existe_no_revienta(self, tmp_path: Path) -> None:
        lectura = leer_plantilla(tmp_path / "fantasma.xlsx")
        assert not lectura.valida
        assert lectura.incidencias[0].mensaje == "no existe"

    def test_leer_o_fallar_levanta_con_todas(self, plantilla_llena: Path) -> None:
        libro = load_workbook(plantilla_llena)
        fila = _escribir(libro["Mina Alfa"], "Mineral extraído", [0.0] * 3)
        libro["Mina Alfa"].cell(row=fila, column=3, value="x")
        libro.save(plantilla_llena)

        with pytest.raises(ErrorDePlantilla, match="incidencia"):
            leer_o_fallar(plantilla_llena)


class TestLibroDeProduccion:
    """El libro que sube el usuario: pestañas de proyecto y nada mas.

    No identifica el caso. El archivo llega desde un caso ya abierto en la
    plataforma, y cada pestaña se asocia a una unidad por su posicion.
    """

    @pytest.fixture
    def libro_de_produccion(self, tmp_path: Path) -> Path:
        generador = _generador()
        ruta = tmp_path / "produccion.xlsx"
        libro = generador.Workbook()  # type: ignore[attr-defined]
        libro.remove(libro.active)
        for nombre in ("Proyecto X", "Proyecto Y"):
            generador.hoja_produccion_de_unidad(libro, nombre, 2027, 45)  # type: ignore[attr-defined]
        generador.hoja_instrucciones(libro, [])  # type: ignore[attr-defined]
        libro.save(ruta)
        return ruta

    def test_no_necesita_hoja_de_caso(self, libro_de_produccion: Path) -> None:
        lectura = leer_produccion(libro_de_produccion)
        assert lectura.valida, [str(i) for i in lectura.incidencias]

    def test_deduce_un_horizonte_largo_de_la_fila_de_anos(self, libro_de_produccion: Path) -> None:
        # Nada fija el numero de anos de antemano: se cuenta la fila de anos, de
        # modo que un proyecto de vida larga no exige tocar el lector.
        horizonte = leer_produccion(libro_de_produccion).horizonte
        assert horizonte is not None
        assert horizonte.primer_ano == 2027
        assert horizonte.anos == 45

    def test_las_pestanas_llegan_en_orden_y_sin_identidad(self, libro_de_produccion: Path) -> None:
        bloques = leer_produccion(libro_de_produccion).bloques
        assert [b.orden for b in bloques] == [1, 2]
        assert [b.hoja for b in bloques] == ["Proyecto X", "Proyecto Y"]

    def test_todos_los_proyectos_tienen_la_misma_estructura(
        self, libro_de_produccion: Path
    ) -> None:
        # Es una plantilla general y no se adapta al proyecto: uno sin
        # preconcentracion deja esas filas en cero.
        libro = load_workbook(libro_de_produccion)
        etiquetas = {
            hoja: [
                libro[hoja].cell(row=f, column=1).value
                for f in range(5, libro[hoja].max_row + 1)
                if libro[hoja].cell(row=f, column=2).value
            ]
            for hoja in ("Proyecto X", "Proyecto Y")
        }
        assert etiquetas["Proyecto X"] == etiquetas["Proyecto Y"]
        assert etiquetas["Proyecto X"] == [f.etiqueta for f in FILAS_CON_DATO]

    def test_una_fila_de_anos_con_saltos_se_reporta(self, libro_de_produccion: Path) -> None:
        # Un salto desplaza todas las series a partir de ahi, y el
        # desplazamiento no deja rastro en el resultado.
        libro = load_workbook(libro_de_produccion)
        libro["Proyecto X"].cell(row=3, column=10, value=2999)
        libro.save(libro_de_produccion)

        lectura = leer_produccion(libro_de_produccion)
        assert not lectura.valida
        assert any("consecutivos" in i.mensaje for i in lectura.incidencias)

    def test_un_libro_sin_pestanas_de_proyecto_se_reporta(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        ruta = tmp_path / "vacio.xlsx"
        libro = Workbook()
        libro.create_sheet("Leeme")
        del libro["Sheet"]
        libro.save(ruta)

        lectura = leer_produccion(ruta)
        assert not lectura.valida
        assert "ninguna pestana de proyecto" in lectura.incidencias[0].mensaje


class TestAsociacionPorOrden:
    def test_empareja_pestana_con_unidad_por_posicion(self) -> None:
        assert asociar_por_orden(["Proyecto X", "Proyecto Y"], ["SR", "B2"]) == {
            "SR": "Proyecto X",
            "B2": "Proyecto Y",
        }

    def test_el_nombre_de_la_pestana_no_decide(self) -> None:
        # Quien llena el archivo rotula como quiera. Si el nombre mandara, una
        # pestana llamada "B2" acabaria en la unidad equivocada.
        assert asociar_por_orden(["B2", "SR"], ["SR", "B2"])["SR"] == "B2"

    def test_si_sobran_o_faltan_pestanas_no_adivina(self) -> None:
        with pytest.raises(ErrorDeAsociacion, match="sobra o falta"):
            asociar_por_orden(["Proyecto X"], ["SR", "B2"])

    def test_el_selector_provisional_no_incluye_la_refineria(self) -> None:
        # Pisco no se carga: se calcula a partir de lo que le entregan las minas.
        assert UNIDADES_PROVISIONALES == ("SR", "B2", "NZ", "SRP", "SD")


class TestSinonimos:
    def test_entiende_las_abreviaturas_del_libro(self) -> None:
        assert canonizar("LT")[0] == "linea de transmision"
        assert canonizar("Pta Subproductos")[0] == "planta de subproductos"

    def test_quita_el_nombre_de_la_unidad_pegado(self) -> None:
        # El libro pega el nombre de la unidad a la etiqueta. Y `Tratamiento de
        # Relaves B2` es una linea de costo operativo, no de produccion: hasta
        # el 01/09/2026 el sinonimo la mandaba a un concepto de tonelaje.
        assert canonizar("Tratamiento de Relaves B2", unidades=["B2"])[0] == "relavera"

    def test_separa_el_metal_del_concepto(self) -> None:
        assert canonizar("Ley de cabeza de Sn") == ("ley de cabeza", "Sn")
        assert canonizar("Recuperacion Cu") == ("recuperacion", "Cu")


def test_la_lectura_declara_si_es_valida(plantilla_llena: Path) -> None:
    assert Lectura(None, ()).valida is False
    assert leer_plantilla(plantilla_llena).valida is True


def _libro_de_supuestos(generador: object, carpeta: Path, nombres: list[str], anos: int) -> Path:
    """La segunda plantilla: sin ella la produccion se lee y no se puede vender."""
    ruta = carpeta / "supuestos.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    generador.hoja_supuestos(libro, 2030, anos)  # type: ignore[attr-defined]
    for nombre in nombres:
        generador.hoja_supuestos_de_unidad(libro, nombre, 2030, anos)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    _escribir(libro["Comunes"], "Capacidad Maxima de la Refineria", [100_000.0] * anos)
    _escribir(libro["Comunes"], "Premio Sn", [0.0] * anos)
    for nombre in nombres:
        _escribir(libro[nombre], "Recuperacion de Sn en la refineria", [0.95] * anos)
    libro.save(ruta)
    return ruta


def _libro_de_precios(
    generador: object, carpeta: Path, anos: int, precio: float = 30_000.0
) -> Path:
    ruta = carpeta / "precios.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    generador.hoja_comite_de_precios(libro, 2030, anos)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    hoja = libro["Comite de Precios"]
    hoja["B2"] = "Comite de prueba"
    hoja["D2"] = "01/01/2030"
    _escribir(hoja, "Sn", [precio] * anos)
    libro.save(ruta)
    return ruta


class TestUnProyectoCualquieraLlegaAlIndicador:
    """La plantilla no sabe de que proyecto se trata, y el motor tampoco.

    Es la propiedad que sostiene el alcance: la plataforma evalua un proyecto
    que hoy no existe con la misma estructura y el mismo motor que las unidades
    en marcha. Aqui se ejercita de punta a punta -generar, llenar, leer, aplicar
    los supuestos y calcular- sobre nombres que no son de nadie.
    """

    def test_tres_proyectos_inventados_recorren_la_cadena_hasta_el_npv(
        self, tmp_path: Path
    ) -> None:
        generador = _generador()
        anos = 3
        nombres = ["Proyecto X", "Proyecto Y", "Proyecto Z"]
        unidades = generador._encadenar(  # type: ignore[attr-defined]
            [
                generador.Unidad.desde_texto(f"{n}:mina:Sn")  # type: ignore[attr-defined]
                for n in nombres
            ]
            + [generador.Unidad.desde_texto("Refineria:refineria:Sn")]  # type: ignore[attr-defined]
        )
        ruta = tmp_path / "caso.xlsx"
        libro = generador.Workbook()  # type: ignore[attr-defined]
        generador.hoja_caso(libro, unidades, 2030, anos)  # type: ignore[attr-defined]
        for nombre in nombres:
            generador.hoja_produccion_de_unidad(libro, nombre, 2030, anos)  # type: ignore[attr-defined]
        generador.hoja_precios(libro, unidades, 2030, anos)  # type: ignore[attr-defined]
        libro.save(ruta)

        libro = load_workbook(ruta)
        libro["Caso"]["B3"] = "Cartera inventada"
        libro["Caso"]["B9"] = "CP-PRUEBA"
        for nombre in nombres:
            _llenar_cadena(libro[nombre])
        _escribir(libro["Precios"], "Precio, escenario Base", [30_000.0] * anos)
        _escribir(libro["Precios"], "Factor de metal pagable", [0.9] * anos)
        libro.save(ruta)

        caso = leer_o_fallar(ruta)
        assert [u.nombre for u in caso.unidades] == [*nombres, "Refineria"]

        supuestos = leer_supuestos(_libro_de_supuestos(generador, tmp_path, nombres, anos))
        comite = leer_comite_de_precios(_libro_de_precios(generador, tmp_path, anos))
        assert supuestos.valida, [str(i) for i in supuestos.incidencias]
        assert comite.valida, [str(i) for i in comite.incidencias]
        assert supuestos.supuestos is not None

        corrida = calcular(aplicar_supuestos(caso, supuestos.supuestos, comite.comite), MAESTROS)

        # Los tres aportan, y el motor los lleva sin agrupar.
        for nombre in nombres:
            assert corrida.mineral_tratado_por_unidad[nombre] == (0.0, 1_000.0, 1_000.0)
        # La cadena llega al final: hay venta, hay flujo y hay indicador.
        assert sum(corrida.ventas) > 0.0, "la produccion no llego a venderse"
        assert math.isfinite(corrida.indicadores.npv)
        assert corrida.indicadores.npv != 0.0
        # Y el horizonte es el que declara la plantilla, no uno fijado en el codigo.
        assert caso.horizonte.primer_ano == 2030
        assert len(corrida.flujo.flujo_economico) == anos

    def test_el_indicador_lo_mueve_el_comite_y_no_una_constante(self, tmp_path: Path) -> None:
        # Dos cosas a la vez. Que el NPV se mueva al cambiar el precio prueba que
        # sale de la plantilla y no del codigo; que se mueva con el **comite** y
        # no con la hoja del caso prueba que el dato maestro manda, que es la
        # regla que impide cambiar un precio aprobado sin que nadie lo advierta.
        def npv_con(precio: float) -> float:
            generador = _generador()
            carpeta = tmp_path / f"p{int(precio)}"
            carpeta.mkdir()
            unidades = generador._encadenar(  # type: ignore[attr-defined]
                [
                    generador.Unidad.desde_texto("Proyecto X:mina:Sn"),  # type: ignore[attr-defined]
                    generador.Unidad.desde_texto("Refineria:refineria:Sn"),  # type: ignore[attr-defined]
                ]
            )
            ruta = carpeta / "caso.xlsx"
            libro = generador.Workbook()  # type: ignore[attr-defined]
            generador.hoja_caso(libro, unidades, 2030, 3)  # type: ignore[attr-defined]
            generador.hoja_produccion_de_unidad(libro, "Proyecto X", 2030, 3)  # type: ignore[attr-defined]
            generador.hoja_precios(libro, unidades, 2030, 3)  # type: ignore[attr-defined]
            libro.save(ruta)
            libro = load_workbook(ruta)
            libro["Caso"]["B3"] = "Un proyecto"
            libro["Caso"]["B9"] = "CP-PRUEBA"
            _llenar_cadena(libro["Proyecto X"])
            # La hoja del caso queda fija: la que se mueve es la del comite.
            _escribir(libro["Precios"], "Precio, escenario Base", [30_000.0] * 3)
            _escribir(libro["Precios"], "Factor de metal pagable", [0.9] * 3)
            libro.save(ruta)

            supuestos = leer_supuestos(_libro_de_supuestos(generador, carpeta, ["Proyecto X"], 3))
            comite = leer_comite_de_precios(_libro_de_precios(generador, carpeta, 3, precio))
            assert supuestos.supuestos is not None
            caso = aplicar_supuestos(leer_o_fallar(ruta), supuestos.supuestos, comite.comite)
            return calcular(caso, MAESTROS).indicadores.npv

        barato, caro = npv_con(30_000.0), npv_con(60_000.0)
        assert caro > barato, "el NPV no reacciona al precio del comite"
        assert barato != 0.0
