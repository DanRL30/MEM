"""Pruebas de las dos plantillas de supuestos.

Son dos archivos separados y con dueños distintos: el comité de precios lo
aprueba y lo sube Finanzas, y el resto de los supuestos son del caso. Estas
pruebas fijan esa separación —ninguna de las dos lee la otra— y el contrato de
cada una con su generador.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from minsur_engine.caso import (
    Caso,
    DatosComunes,
    DatosMaestros,
    ProduccionDeUnidad,
    TerminosComerciales,
    UnidadProductiva,
)
from minsur_engine.corrida import calcular
from minsur_engine.depreciacion import TasasDeDepreciacion
from minsur_engine.horizonte import Horizonte
from minsur_engine.parametros import ParametrosCorporativos
from minsur_engine.tributos import EscalaProgresiva, Tramo
from minsur_ingest.plantilla import leer_comite_de_precios, leer_supuestos
from minsur_ingest.supuestos import CON_DATO_DE_SUPUESTOS, CON_DATO_POR_UNIDAD, aplicar

MAESTROS = DatosMaestros(
    parametros=ParametrosCorporativos(
        version_datos_maestros="CP-PRUEBA",
        tasa_descuento=0.10,
        participacion_trabajadores=0.08,
        impuesto_renta=0.295,
        regalia_minima=0.01,
        osinergmin=0.0014,
        oefa=0.001,
        fondo_jubilacion_minera=0.005,
    ),
    tasas_tributarias=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    tasas_financieras=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    escala_regalia=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.01),)),
    escala_iem=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.0),)),
)

RAIZ = Path(__file__).resolve().parents[3]


def _generador() -> object:
    ruta = RAIZ / "scripts" / "generar_plantilla_inputs.py"
    especificacion = importlib.util.spec_from_file_location("generador", ruta)
    assert especificacion is not None and especificacion.loader is not None
    modulo = importlib.util.module_from_spec(especificacion)
    sys.modules["generador"] = modulo
    especificacion.loader.exec_module(modulo)
    return modulo


def _llenar(hoja: object, valores: list[float], anos: int = 3) -> None:
    """Llena por posición las filas con unidad de medida."""
    fila, i = 5, 0
    while i < len(valores):
        if hoja.cell(row=fila, column=2).value:  # type: ignore[attr-defined]
            for ano in range(anos):
                hoja.cell(row=fila, column=3 + ano, value=valores[i])  # type: ignore[attr-defined]
            i += 1
        fila += 1


@pytest.fixture
def comite(tmp_path: Path) -> Path:
    generador = _generador()
    ruta = tmp_path / "precios.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    generador.hoja_comite_de_precios(libro, 2027, 3)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    hoja = libro["Comite de Precios"]
    hoja["B2"] = "Comite de Precios Mar-2026"
    hoja["D2"] = "15/03/2026"
    _llenar(hoja, [30_000.0, 9_000.0, 31.0])
    libro.save(ruta)
    return ruta


@pytest.fixture
def supuestos(tmp_path: Path) -> Path:
    generador = _generador()
    ruta = tmp_path / "supuestos.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    generador.hoja_supuestos(libro, 2027, 3)  # type: ignore[attr-defined]
    for nombre in ("Proyecto X", "Proyecto Y"):
        generador.hoja_supuestos_de_unidad(libro, nombre, 2027, 3)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    _llenar(libro["Comunes"], [float(i) for i in range(len(CON_DATO_DE_SUPUESTOS))])
    _llenar(libro["Proyecto X"], [30_000.0, 37_000.0, 95.0])
    _llenar(libro["Proyecto Y"], [18_000.0, 52_000.0, 70.0])
    libro.save(ruta)
    return ruta


class TestComiteDePrecios:
    def test_lee_los_tres_metales_y_su_identificacion(self, comite: Path) -> None:
        lectura = leer_comite_de_precios(comite)
        assert lectura.valida, [str(i) for i in lectura.incidencias]
        assert lectura.comite is not None
        assert lectura.comite.nombre == "Comite de Precios Mar-2026"
        assert lectura.comite.aprobado_el == "15/03/2026"
        assert lectura.comite.sn == (30_000.0, 30_000.0, 30_000.0)
        assert lectura.comite.cu == (9_000.0, 9_000.0, 9_000.0)
        assert lectura.comite.ag == (31.0, 31.0, 31.0)

    def test_un_comite_sin_nombre_se_rechaza(self, comite: Path) -> None:
        # Una corrida registra que comite uso, no "el vigente". Sin nombre no
        # hay a que referirse, y el caso quedaria sin poder reproducirse.
        libro = load_workbook(comite)
        libro["Comite de Precios"]["B2"] = None
        libro.save(comite)

        lectura = leer_comite_de_precios(comite)
        assert not lectura.valida
        assert any("no esta identificado" in i.mensaje for i in lectura.incidencias)

    def test_deduce_el_horizonte_de_la_fila_de_anos(self, comite: Path) -> None:
        lectura = leer_comite_de_precios(comite)
        assert lectura.comite is not None
        assert lectura.comite.horizonte.primer_ano == 2027
        assert lectura.comite.horizonte.anos == 3

    def test_no_acepta_la_plantilla_de_supuestos(self, supuestos: Path) -> None:
        # Son dos archivos con dueños distintos. Confundirlos dejaria a
        # cualquiera subiendo precios por la puerta del caso.
        lectura = leer_comite_de_precios(supuestos)
        assert not lectura.valida
        assert "No es la plantilla del comite" in lectura.incidencias[0].mensaje


class TestSupuestosDelCaso:
    def test_lee_los_comunes_y_los_de_cada_unidad(self, supuestos: Path) -> None:
        lectura = leer_supuestos(supuestos)
        assert lectura.valida, [str(i) for i in lectura.incidencias]
        assert lectura.supuestos is not None
        assert set(lectura.supuestos.por_unidad) == {"Proyecto X", "Proyecto Y"}
        assert len(lectura.supuestos.comunes) == len(CON_DATO_DE_SUPUESTOS)
        assert len(lectura.supuestos.por_unidad["Proyecto X"]) == len(CON_DATO_POR_UNIDAD)

    def test_los_miles_de_dolares_se_convierten(self, supuestos: Path) -> None:
        # Regla 003: la depreciacion del libro va en k$. La conversion ocurre en
        # la frontera y el motor trabaja siempre en dolares.
        de_x = leer_supuestos(supuestos).supuestos
        assert de_x is not None
        assert de_x.por_unidad["Proyecto X"]["depreciacion_tributaria"] == (
            30_000_000.0,
            30_000_000.0,
            30_000_000.0,
        )

    def test_los_porcentajes_se_convierten(self, supuestos: Path) -> None:
        de_x = leer_supuestos(supuestos).supuestos
        assert de_x is not None
        assert de_x.por_unidad["Proyecto X"]["recuperacion_en_la_refineria"] == pytest.approx(
            (0.95, 0.95, 0.95)
        )

    def test_la_recuperacion_en_la_refineria_va_por_unidad(self, supuestos: Path) -> None:
        # Regla de oro: el libro la agrupa en `SR + B2` y `NZ + SRP`, y aqui
        # cada unidad tiene la suya.
        por_unidad = leer_supuestos(supuestos).supuestos
        assert por_unidad is not None
        assert (
            por_unidad.por_unidad["Proyecto X"]["recuperacion_en_la_refineria"]
            != (por_unidad.por_unidad["Proyecto Y"]["recuperacion_en_la_refineria"])
        )

    def test_una_estructura_alterada_se_reporta(self, supuestos: Path) -> None:
        libro = load_workbook(supuestos)
        libro["Comunes"].cell(row=6, column=1, value="Prima del estano")
        libro.save(supuestos)

        lectura = leer_supuestos(supuestos)
        assert not lectura.valida
        assert any("estructura fija" in i.mensaje for i in lectura.incidencias)

    def test_no_acepta_la_plantilla_del_comite(self, comite: Path) -> None:
        lectura = leer_supuestos(comite)
        assert not lectura.valida
        assert "No es la plantilla de supuestos" in lectura.incidencias[0].mensaje

    def test_un_archivo_que_no_existe_no_revienta(self, tmp_path: Path) -> None:
        assert leer_supuestos(tmp_path / "fantasma.xlsx").incidencias[0].mensaje == "no existe"


class TestAplicarAlCaso:
    """La costura entre las tres plantillas.

    Produccion trae las series de cada unidad; el comite, lo que vale cada
    metal; y los supuestos, como se liquida. Sin este paso la produccion se lee
    y no se puede vender.
    """

    def _caso(self) -> Caso:
        horizonte = Horizonte(primer_ano=2027, anos=3)
        ceros = horizonte.ceros()
        mina = UnidadProductiva(
            nombre="Mina Alfa",
            tipo="mina",
            produccion=ProduccionDeUnidad(
                mineral_tratado=horizonte.serie((0.0, 1_000.0, 1_000.0), nombre="tratado"),
                concentrado_producido=horizonte.serie((0.0, 500.0, 500.0), nombre="concentrado"),
                ley_del_concentrado=horizonte.serie((0.0, 0.40, 0.40), nombre="ley"),
                concentrado_de_cu=horizonte.serie((0.0, 200.0, 200.0), nombre="cu"),
            ),
        )
        refinería = UnidadProductiva(
            nombre="Refineria",
            tipo="refineria",
            produccion=ProduccionDeUnidad(mineral_tratado=ceros),
        )
        return Caso(
            nombre="Caso de prueba",
            horizonte=horizonte,
            unidades=(mina, refinería),
            terminos=TerminosComerciales(
                precio_metal_refinado=ceros,
                premio_metal_refinado=ceros,
                precio_metal_en_concentrado=ceros,
                factor_metal_pagable=ceros,
            ),
            datos_comunes=DatosComunes(gastos_administrativos=ceros),
        )

    def test_el_comite_pone_los_precios_y_los_supuestos_lo_demas(
        self, comite: Path, supuestos: Path
    ) -> None:
        leido = leer_comite_de_precios(comite).comite
        del_caso = leer_supuestos(supuestos).supuestos
        assert leido is not None and del_caso is not None

        caso = aplicar(self._caso(), del_caso, leido)
        assert caso.terminos.precio_metal_refinado == (30_000.0, 30_000.0, 30_000.0)
        assert caso.terminos.concentrado is not None
        assert [m.nombre for m in caso.terminos.concentrado.metales] == ["Cu", "Ag"]
        assert caso.terminos.concentrado.metales[1].en_onzas_troy is True

    def test_la_recuperacion_cuelga_de_la_refineria_y_por_origen(
        self, comite: Path, supuestos: Path
    ) -> None:
        # Regla de oro: es la refinería quien refina, y guarda una recuperacion
        # por unidad de origen en vez de una comun a todas.
        del_caso = leer_supuestos(supuestos).supuestos
        assert del_caso is not None
        caso = aplicar(self._caso(), del_caso, leer_comite_de_precios(comite).comite)
        refineria = caso.refineria
        assert refineria is not None
        assert refineria.recuperacion_en_la_refineria["Mina Alfa"]["Sn"] == pytest.approx(
            (0.95, 0.95, 0.95)
        )

    def test_el_concentrado_de_cobre_llega_a_la_venta(self, comite: Path, supuestos: Path) -> None:
        # Es el tercer camino de ingreso del libro y hasta el 01/09/2026 estaba
        # implementado y sin cablear: se calculaba y nadie lo cobraba.
        del_caso = leer_supuestos(supuestos).supuestos
        assert del_caso is not None
        caso = aplicar(self._caso(), del_caso, leer_comite_de_precios(comite).comite)
        corrida = calcular(caso, MAESTROS)
        assert corrida.concentrado_liquidado_por_unidad["Mina Alfa"][1] != 0.0
        assert corrida.ventas[1] != 0.0

    def test_los_reguladores_del_caso_ganan_a_la_tasa_de_referencia(
        self, comite: Path, supuestos: Path
    ) -> None:
        # MINSUR confirmo el 01/09/2026 que varian los primeros anos porque
        # tienen mejor informacion. Si el caso los declara, mandan.
        del_caso = leer_supuestos(supuestos).supuestos
        assert del_caso is not None
        caso = aplicar(self._caso(), del_caso, leer_comite_de_precios(comite).comite)
        assert caso.datos_comunes.oefa != ()
        assert caso.datos_comunes.osinergmin != ()
