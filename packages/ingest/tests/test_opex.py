"""Pruebas de la plantilla de opex y su lectura.

La prueba que sostiene el paquete es la de ida y vuelta: se genera la plantilla
con el script que la emite, se llena por posición, se lee y se calcula. Generador
y lector son las dos mitades del mismo contrato y esta prueba falla en cuanto una
se mueve sin la otra.

Las cifras son redondas a propósito: el cash cost de la mina suma 300 y el de
la refinería 100, de modo que cualquier desvío se ve a simple vista.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from minsur_engine.cash_cost import (
    GESTION_SOCIAL,
    GESTION_SOCIAL_DEDUCIBLE,
    PLANILLA,
)
from minsur_engine.caso import (
    Caso,
    DatosComunes,
    DatosMaestros,
    ProduccionDeUnidad,
    TerminosComerciales,
    UnidadProductiva,
)
from minsur_engine.corrida import calcular
from minsur_engine.depreciacion import TasasDeDepreciacion
from minsur_engine.horizonte import Horizonte
from minsur_engine.impuestos import EscalaProgresiva, Tramo
from minsur_engine.parametros import ParametrosCorporativos
from minsur_ingest.opex import (
    CON_DATO_DE_CASH_COST,
    CON_DATO_DE_GASTOS,
    CONCEPTOS_LIBRES,
    ErrorDeAsociacion,
    aplicar,
)
from minsur_ingest.plantilla import leer_opex

RAIZ = Path(__file__).resolve().parents[3]

MAESTROS = DatosMaestros(
    parametros=ParametrosCorporativos(
        version_datos_maestros="CP-PRUEBA",
        tasa_descuento=0.10,
        participacion_trabajadores=0.08,
        impuesto_renta=0.295,
        regalia_minima=0.01,
        osinergmin=0.0,
        oefa=0.0,
        fondo_jubilacion_minera=0.005,
    ),
    tasas_tributarias=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    tasas_financieras=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    escala_regalia=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.01),)),
    escala_iem=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.0),)),
)

ANOS = 3


def _generador() -> object:
    ruta = RAIZ / "scripts" / "generar_plantilla_inputs.py"
    especificacion = importlib.util.spec_from_file_location("generador", ruta)
    assert especificacion is not None and especificacion.loader is not None
    modulo = importlib.util.module_from_spec(especificacion)
    sys.modules["generador"] = modulo
    especificacion.loader.exec_module(modulo)
    return modulo


def _escribir(hoja: object, etiqueta: str, valores: list[float]) -> int:
    """Escribe una fila buscándola por su etiqueta, y devuelve su número."""
    for fila in hoja.iter_rows(min_row=5, max_col=2):  # type: ignore[attr-defined]
        celda = fila[0]
        if celda.value and str(celda.value).strip() == etiqueta:
            for i, valor in enumerate(valores):
                hoja.cell(row=celda.row, column=3 + i, value=valor)  # type: ignore[attr-defined]
            return int(celda.row)
    raise AssertionError(f"la plantilla no tiene la fila {etiqueta!r}")


@pytest.fixture
def libro_de_opex(tmp_path: Path) -> Path:
    """Un libro de dos pestañas —una mina y la refinería— con costos y gastos."""
    generador = _generador()
    ruta = tmp_path / "opex.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    for nombre in ("Mina Alfa", "Refineria"):
        generador.hoja_opex_de_unidad(libro, nombre, 2027, ANOS)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    mina = libro["Mina Alfa"]
    _escribir(mina, "Mina", [0.0, 200.0, 200.0])
    _escribir(mina, "Planta Concentradora", [0.0, 100.0, 100.0])
    _escribir(mina, "Gastos administrativos", [0.0, 50.0, 50.0])
    _escribir(mina, "Gestión Social", [0.0, 40.0, 40.0])
    _escribir(mina, "Estudios Pre Factibilidad (Gasto)", [30.0, 0.0, 0.0])
    _escribir(mina, "Estudios Factibilidad (Capitalizable)", [70.0, 0.0, 0.0])
    _escribir(libro["Refineria"], "Fundición", [0.0, 100.0, 100.0])
    libro.save(ruta)
    return ruta


def _caso() -> Caso:
    horizonte = Horizonte(primer_ano=2027, anos=ANOS)
    ceros = horizonte.ceros()
    return Caso(
        nombre="Caso de prueba",
        horizonte=horizonte,
        unidades=(
            UnidadProductiva(
                nombre="Mina Alfa",
                tipo="mina",
                produccion=ProduccionDeUnidad(
                    mineral_tratado=horizonte.serie((0.0, 1_000.0, 1_000.0), nombre="tratado")
                ),
            ),
            UnidadProductiva(
                nombre="Refineria",
                tipo="refineria",
                produccion=ProduccionDeUnidad(mineral_tratado=ceros),
            ),
        ),
        terminos=TerminosComerciales(
            precio_metal_refinado=ceros,
            premio_metal_refinado=ceros,
            precio_metal_en_concentrado=ceros,
            factor_metal_pagable=ceros,
        ),
        datos_comunes=DatosComunes(gastos_administrativos=ceros),
    )


def _con_opex(ruta: Path) -> Caso:
    lectura = leer_opex(ruta)
    assert lectura.valida, [str(i) for i in lectura.incidencias]
    return aplicar(_caso(), [b.opex for b in lectura.bloques])


class TestIdaYVuelta:
    def test_la_plantilla_llena_se_lee(self, libro_de_opex: Path) -> None:
        lectura = leer_opex(libro_de_opex)
        assert lectura.valida, [str(i) for i in lectura.incidencias]
        assert [b.hoja for b in lectura.bloques] == ["Mina Alfa", "Refineria"]
        assert lectura.horizonte is not None
        assert lectura.horizonte.primer_ano == 2027

    def test_las_dos_pestanas_tienen_la_misma_estructura(self, libro_de_opex: Path) -> None:
        # Es lo que permite cargar un proyecto que hoy no existe en el libro: lo
        # que no aplica va en cero y la plataforma no lo muestra.
        libro = load_workbook(libro_de_opex)
        esperadas = [f.etiqueta for f in CON_DATO_DE_CASH_COST + CON_DATO_DE_GASTOS]
        for nombre in ("Mina Alfa", "Refineria"):
            hoja = libro[nombre]
            leidas = [
                str(fila[0].value).strip()
                for fila in hoja.iter_rows(min_row=5, max_col=2)
                if fila[1].value and fila[0].value
            ]
            assert leidas == esperadas

    def test_los_costos_y_los_gastos_llegan_separados(self, libro_de_opex: Path) -> None:
        mina = _con_opex(libro_de_opex).unidades[0]
        assert mina.costos["Mina"] == (0.0, 200_000.0, 200_000.0)
        assert mina.gastos["Gastos administrativos"] == (0.0, 50_000.0, 50_000.0)
        assert "Gastos administrativos" not in mina.costos

    def test_la_refineria_tambien_lleva_pestana(self, libro_de_opex: Path) -> None:
        # Su produccion es resultado, pero su costo es dato: por eso el libro de
        # opex trae una pestana mas que el de produccion.
        refineria = _con_opex(libro_de_opex).unidades[1]
        assert refineria.es_refineria
        assert refineria.costos["Fundición"] == (0.0, 100_000.0, 100_000.0)

    def test_el_caso_leido_calcula(self, libro_de_opex: Path) -> None:
        corrida = calcular(_con_opex(libro_de_opex), MAESTROS)
        assert corrida.cash_cost == (0.0, 400_000.0, 400_000.0)
        assert corrida.cash_cost_por_unidad["Mina Alfa"] == (0.0, 300_000.0, 300_000.0)
        assert corrida.cash_cost_por_unidad["Refineria"] == (0.0, 100_000.0, 100_000.0)


class TestConversionDeEscalas:
    def test_los_miles_de_dolares_se_convierten(self, libro_de_opex: Path) -> None:
        # Regla 003: el libro lleva el bloque entero en miles de dolares y el
        # motor trabaja en dolares. La conversion ocurre en la ingesta.
        assert _con_opex(libro_de_opex).unidades[0].costos["Planta Concentradora"] == (
            0.0,
            100_000.0,
            100_000.0,
        )

    def test_la_unidad_declarada_manda_sobre_la_de_la_plantilla(self, libro_de_opex: Path) -> None:
        libro = load_workbook(libro_de_opex)
        hoja = libro["Mina Alfa"]
        fila = _escribir(hoja, "Mina", [0.0, 200_000.0, 200_000.0])
        hoja.cell(row=fila, column=2, value="US$")
        libro.save(libro_de_opex)
        assert _con_opex(libro_de_opex).unidades[0].costos["Mina"] == (
            0.0,
            200_000.0,
            200_000.0,
        )


class TestLaCola:
    def test_un_concepto_propio_entra_y_suma_al_total(self, libro_de_opex: Path) -> None:
        # Acuerdo 6 de la minuta del 27/08/2026: la lista es extensible y lo
        # anadido solo afecta al total.
        libro = load_workbook(libro_de_opex)
        hoja = libro["Mina Alfa"]
        fila = _escribir(hoja, CON_DATO_DE_CASH_COST[-1].etiqueta, [0.0] * ANOS) + 2
        hoja.cell(row=fila, column=1, value="Servidumbre de paso del ferrocarril")
        for i, valor in enumerate([0.0, 25.0, 25.0]):
            hoja.cell(row=fila, column=3 + i, value=valor)
        libro.save(libro_de_opex)

        caso = _con_opex(libro_de_opex)
        mina = caso.unidades[0]
        assert mina.costos["Servidumbre de paso del ferrocarril"] == (0.0, 25_000.0, 25_000.0)
        assert calcular(caso, MAESTROS).cash_cost_por_unidad["Mina Alfa"] == (
            0.0,
            325_000.0,
            325_000.0,
        )

    def test_mas_conceptos_de_los_que_admite_la_cola_se_reportan(self, libro_de_opex: Path) -> None:
        # Tragarlos convertiria un gasto mal escrito en un costo con su nombre.
        libro = load_workbook(libro_de_opex)
        hoja = libro["Mina Alfa"]
        primera = _escribir(hoja, CON_DATO_DE_CASH_COST[-1].etiqueta, [0.0] * ANOS) + 1
        for salto in range(CONCEPTOS_LIBRES + 1):
            hoja.cell(row=primera + salto, column=1, value=f"Concepto propio {salto}")
            hoja.cell(row=primera + salto, column=2, value="$k")
        libro.save(libro_de_opex)

        lectura = leer_opex(libro_de_opex)
        assert not lectura.valida
        assert any("cola de conceptos propios" in i.mensaje for i in lectura.incidencias)

    def test_importes_sin_concepto_se_reportan(self, libro_de_opex: Path) -> None:
        # Es el fallo silencioso que la plantilla de produccion enseno a temer:
        # un costo que nadie puede atribuir y que desaparece del total.
        libro = load_workbook(libro_de_opex)
        hoja = libro["Mina Alfa"]
        fila = _escribir(hoja, CON_DATO_DE_CASH_COST[-1].etiqueta, [0.0] * ANOS) + 2
        hoja.cell(row=fila, column=3, value=90.0)
        libro.save(libro_de_opex)

        lectura = leer_opex(libro_de_opex)
        assert not lectura.valida
        incidencia = next(i for i in lectura.incidencias if "sin concepto" in i.mensaje)
        assert incidencia.hoja == "Mina Alfa"
        assert incidencia.celda == f"A{fila}"


class TestGastosDerivados:
    def test_la_plantilla_no_pide_lo_que_el_libro_deriva(self) -> None:
        etiquetas = [f.etiqueta for f in CON_DATO_DE_GASTOS]
        assert PLANILLA not in etiquetas
        assert GESTION_SOCIAL_DEDUCIBLE not in etiquetas

    def test_la_planilla_sale_del_cash_cost_de_la_unidad(self, libro_de_opex: Path) -> None:
        caso = _con_opex(libro_de_opex)
        comunes = caso.datos_comunes
        con_tasa = caso.__class__(
            nombre=caso.nombre,
            horizonte=caso.horizonte,
            unidades=caso.unidades,
            terminos=caso.terminos,
            datos_comunes=DatosComunes(
                gastos_administrativos=comunes.gastos_administrativos,
                planilla_sobre_cash_cost=caso.horizonte.serie((0.1,) * ANOS, nombre="planilla"),
            ),
        )
        gastos = calcular(con_tasa, MAESTROS).gastos_por_unidad["Mina Alfa"]
        assert gastos[PLANILLA] == pytest.approx((0.0, 30_000.0, 30_000.0))

    def test_la_gestion_social_es_deducible_entera_si_no_se_declara_fraccion(
        self, libro_de_opex: Path
    ) -> None:
        gastos = calcular(_con_opex(libro_de_opex), MAESTROS).gastos_por_unidad["Mina Alfa"]
        assert gastos[GESTION_SOCIAL_DEDUCIBLE] == gastos[GESTION_SOCIAL]

    def test_una_fraccion_declarada_rebaja_solo_la_base_imponible(
        self, libro_de_opex: Path
    ) -> None:
        from dataclasses import replace

        caso = _con_opex(libro_de_opex)
        mitad = caso.horizonte.serie((0.5,) * ANOS, nombre="fraccion")
        con_fraccion = replace(
            caso,
            unidades=(
                replace(caso.unidades[0], fraccion_gestion_social_deducible=mitad),
                caso.unidades[1],
            ),
        )
        gastos = calcular(con_fraccion, MAESTROS).gastos_por_unidad["Mina Alfa"]
        assert gastos[GESTION_SOCIAL] == (0.0, 40_000.0, 40_000.0)
        assert gastos[GESTION_SOCIAL_DEDUCIBLE] == (0.0, 20_000.0, 20_000.0)

    def test_el_estudio_capitalizable_sale_de_caja_y_no_rebaja_la_base(
        self, libro_de_opex: Path
    ) -> None:
        # El libro lo lleva a `Depreciacion` en vez de deducirlo. Aqui sale de
        # caja con los estudios y queda fuera del gasto deducible.
        corrida = calcular(_con_opex(libro_de_opex), MAESTROS)
        # Los 30 del estudio de gasto y los 70 del capitalizable salen de caja.
        assert corrida.flujo.flujo_de_inversiones[0] == pytest.approx(-100_000.0)
        # De la base imponible solo se descuentan los 30.
        assert corrida.impuestos.por_ano[0].utilidad_imponible == pytest.approx(-30_000.0)

    def test_el_estudio_capitalizable_se_deprecia(self, libro_de_opex: Path) -> None:
        # Capitalizar es diferir, no perder: sale de caja el primer ano y rebaja
        # la base a lo largo de la vida del activo. El libro lo deprecia al 5 %,
        # y la puerta de la regla 013 lo retrasa hasta el primer ano con
        # produccion, que aqui es el segundo.
        corrida = calcular(_con_opex(libro_de_opex), MAESTROS)
        componentes = corrida.depreciacion_tributaria_por_componente["Mina Alfa"]
        assert componentes["Estudios capitalizables"][1] == pytest.approx(3_500.0)


class TestIncidencias:
    def test_una_estructura_alterada_se_reporta_y_no_se_adivina(self, libro_de_opex: Path) -> None:
        libro = load_workbook(libro_de_opex)
        libro["Mina Alfa"].cell(row=8, column=1, value="Mina subterranea")
        libro.save(libro_de_opex)

        lectura = leer_opex(libro_de_opex)
        assert not lectura.valida
        assert any("estructura fija" in i.mensaje for i in lectura.incidencias)

    def test_una_celda_con_texto_se_reporta_con_su_ubicacion(self, libro_de_opex: Path) -> None:
        libro = load_workbook(libro_de_opex)
        fila = _escribir(libro["Mina Alfa"], "Mina", [0.0] * ANOS)
        libro["Mina Alfa"].cell(row=fila, column=4, value="doscientos")
        libro.save(libro_de_opex)

        lectura = leer_opex(libro_de_opex)
        assert not lectura.valida
        incidencia = lectura.incidencias[0]
        assert incidencia.hoja == "Mina Alfa"
        assert incidencia.celda == f"D{fila}"

    def test_un_bloque_incompleto_se_reporta(self, libro_de_opex: Path) -> None:
        libro = load_workbook(libro_de_opex)
        hoja = libro["Mina Alfa"]
        hoja.delete_rows(_escribir(hoja, "Exploraciones", [0.0] * ANOS))
        libro.save(libro_de_opex)

        lectura = leer_opex(libro_de_opex)
        assert not lectura.valida

    def test_un_archivo_que_no_existe_no_revienta(self, tmp_path: Path) -> None:
        lectura = leer_opex(tmp_path / "fantasma.xlsx")
        assert not lectura.valida
        assert lectura.incidencias[0].mensaje == "no existe"


class TestAsociacionPorOrden:
    def test_las_pestanas_se_asocian_por_posicion(self, libro_de_opex: Path) -> None:
        caso = _con_opex(libro_de_opex)
        assert caso.unidades[0].costos["Mina"] == (0.0, 200_000.0, 200_000.0)
        assert caso.unidades[1].costos["Mina"] == (0.0, 0.0, 0.0)

    def test_una_pestana_de_mas_o_de_menos_se_rechaza(self, libro_de_opex: Path) -> None:
        lectura = leer_opex(libro_de_opex)
        with pytest.raises(ErrorDeAsociacion, match="sobra o falta"):
            aplicar(_caso(), [b.opex for b in lectura.bloques][:1])


@pytest.fixture
def libro_con_servidumbre(tmp_path: Path) -> Path:
    """El mismo libro, con servidumbre y predios cargados por separado.

    Son las filas `InputsOpex!176` y `!177`, que el libro lleva a sitios
    distintos: los predios al flujo de inversiones y la servidumbre a la base
    imponible, a la bolsa de egresos y al flujo operativo.
    """
    generador = _generador()
    ruta = tmp_path / "opex-servidumbre.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    for nombre in ("Mina Alfa", "Refineria"):
        generador.hoja_opex_de_unidad(libro, nombre, 2027, ANOS)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    mina = libro["Mina Alfa"]
    _escribir(mina, "Mina", [0.0, 200.0, 200.0])
    _escribir(mina, "Planta Concentradora", [0.0, 100.0, 100.0])
    _escribir(mina, "Servidumbres y usufructos", [0.0, 7.0, 7.0])
    _escribir(mina, "Predios", [11.0, 0.0, 0.0])
    libro.save(ruta)
    return ruta


class TestLaServidumbreVaDondeElLibroLaPone:
    """`InputsOpex!177` no es un predio, y el libro no la trata como tal.

    Hasta el 02/09/2026 la plataforma sumaba las dos filas en una sola linea de
    inversion, y la servidumbre no rebajaba la base imponible. Es la regla
    `059`, y la decision de alinearse al modelo la tomo el Project Manager el
    mismo dia: donde el libro y la plataforma difieran, manda el libro.
    """

    def test_la_servidumbre_rebaja_la_base_imponible_y_los_predios_no(
        self, libro_con_servidumbre: Path
    ) -> None:
        # `Impuestos!16 = -Otros!49 - Otros!50`, y `Otros!50 = InputsOpex!177`.
        # Los predios no aparecen en ninguna de las dos bases.
        corrida = calcular(_con_opex(libro_con_servidumbre), MAESTROS)
        renta = corrida.impuestos.renta
        assert renta.otros_gastos == (0.0, -7_000.0, -7_000.0)
        # El primer ejercicio solo carga predios: la base no se mueve por ellos.
        assert renta.otros_gastos[0] == 0.0

    def test_la_servidumbre_entra_en_las_dos_bases_por_igual(
        self, libro_con_servidumbre: Path
    ) -> None:
        corrida = calcular(_con_opex(libro_con_servidumbre), MAESTROS)
        assert corrida.impuestos.renta.otros_gastos == corrida.impuestos.regalias.otros_gastos

    def test_la_bolsa_lleva_la_servidumbre_y_no_los_predios(
        self, libro_con_servidumbre: Path
    ) -> None:
        # `Otros!44:54` incluye la fila 50, `Servidumbre`, y no la 93,
        # `Compra de Predios`, que vive fuera del bloque.
        corrida = calcular(_con_opex(libro_con_servidumbre), MAESTROS)
        assert corrida.bolsa_de_egresos[0] == pytest.approx(0.0)
        assert corrida.bolsa_de_egresos[1] == pytest.approx(300_000.0 + 7_000.0)

    def test_los_predios_siguen_en_el_flujo_de_inversiones(
        self, libro_con_servidumbre: Path
    ) -> None:
        # `Otros!93 = -InputsOpex!176`. Es lo unico que queda de aquella linea.
        corrida = calcular(_con_opex(libro_con_servidumbre), MAESTROS)
        assert corrida.flujo.flujo_de_inversiones[0] == pytest.approx(-11_000.0)

    def test_la_servidumbre_del_flujo_espera_al_ejercicio_con_gasto(
        self, libro_con_servidumbre: Path
    ) -> None:
        # `Otros!30` multiplica la fila por `FC NZ!8`, `Periodo con gastos`. En
        # el primer ejercicio no hay opex, de modo que ahi no entra al flujo
        # operativo aunque si rebaje la base imponible, que no lleva bandera.
        from dataclasses import replace

        caso = _con_opex(libro_con_servidumbre)
        adelantada = caso.horizonte.serie((5.0, 7.0, 7.0), nombre="servidumbre")
        mina = caso.unidades[0]
        con_gasto_previo = replace(
            caso,
            unidades=(
                replace(mina, gastos={**mina.gastos, "Servidumbres y usufructos": adelantada}),
                caso.unidades[1],
            ),
        )
        corrida = calcular(con_gasto_previo, MAESTROS)
        # La base imponible la descuenta el primer ano.
        assert corrida.impuestos.renta.otros_gastos[0] == pytest.approx(-5.0)
        # El flujo operativo no, porque ese ejercicio no tiene cash cost.
        assert corrida.cash_cost[0] == 0.0
        assert corrida.flujo.flujo_operativo[0] == pytest.approx(0.0)
