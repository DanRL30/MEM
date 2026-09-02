"""Pruebas de la plantilla de capital y su lectura.

La prueba que sostiene el paquete es la de ida y vuelta: se genera la plantilla
con el script que la emite, se llena por posición, se lee y se calcula.

El caso tiene las tres clases de unidad —una mina, una relavera de depósito y la
refinería— porque lo que distingue a esta plantilla de las otras dos es
justamente que las tres llevan pestaña. Las cifras son redondas: la mina invierte
1 000 en maquinaria el primer año, 200 en cómputo el segundo, y el tercero 400 en
edificaciones y 500 de cierre; el depósito, 300 en edificaciones.
"""

from __future__ import annotations

import importlib.util
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from minsur_engine.caso import (
    Caso,
    DatosComunes,
    DatosMaestros,
    ProduccionDeUnidad,
    TerminosComerciales,
    UnidadProductiva,
)
from minsur_engine.corrida import calcular
from minsur_engine.depreciacion import TasasDeDepreciacion
from minsur_engine.horizonte import Horizonte
from minsur_engine.parametros import ParametrosCorporativos
from minsur_engine.tributos import EscalaProgresiva, Tramo
from minsur_ingest.capex import CON_DATO_DE_CAPEX, ErrorDeAsociacion, aplicar
from minsur_ingest.plantilla import leer_capex

RAIZ = Path(__file__).resolve().parents[3]

MAESTROS = DatosMaestros(
    parametros=ParametrosCorporativos(
        version_datos_maestros="CP-PRUEBA",
        tasa_descuento=0.10,
        participacion_trabajadores=0.08,
        impuesto_renta=0.295,
        regalia_minima=0.01,
        osinergmin=0.0,
        oefa=0.0,
        fondo_jubilacion_minera=0.005,
    ),
    tasas_tributarias=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    tasas_financieras=TasasDeDepreciacion(maquinaria=0.5, instalaciones=0.1, edificaciones=0.05),
    escala_regalia=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.01),)),
    escala_iem=EscalaProgresiva(tramos=(Tramo(0.0, 10.0, 0.0),)),
)

ANOS = 3
PESTANAS = ("Mina Alfa", "Relavera B4", "Refineria")


def _generador() -> object:
    ruta = RAIZ / "scripts" / "generar_plantilla_inputs.py"
    especificacion = importlib.util.spec_from_file_location("generador", ruta)
    assert especificacion is not None and especificacion.loader is not None
    modulo = importlib.util.module_from_spec(especificacion)
    sys.modules["generador"] = modulo
    especificacion.loader.exec_module(modulo)
    return modulo


def _escribir(hoja: object, etiqueta: str, valores: list[float]) -> int:
    """Escribe una fila buscándola por su etiqueta, y devuelve su número."""
    for fila in hoja.iter_rows(min_row=5, max_col=2):  # type: ignore[attr-defined]
        celda = fila[0]
        if celda.value and str(celda.value).strip() == etiqueta:
            for i, valor in enumerate(valores):
                hoja.cell(row=celda.row, column=3 + i, value=valor)  # type: ignore[attr-defined]
            return int(celda.row)
    raise AssertionError(f"la plantilla no tiene la fila {etiqueta!r}")


@pytest.fixture
def libro_de_capex(tmp_path: Path) -> Path:
    """Un libro de tres pestañas: la mina, el depósito y la refinería."""
    generador = _generador()
    ruta = tmp_path / "capex.xlsx"
    libro = generador.Workbook()  # type: ignore[attr-defined]
    libro.remove(libro.active)
    for nombre in PESTANAS:
        generador.hoja_capex_de_unidad(libro, nombre, 2027, ANOS)  # type: ignore[attr-defined]
    libro.save(ruta)

    libro = load_workbook(ruta)
    mina = libro["Mina Alfa"]
    _escribir(mina, "Maquinaria, equipos y vehículos", [1_000.0, 0.0, 0.0])
    _escribir(mina, "Equipos de cómputo", [0.0, 200.0, 0.0])
    _escribir(mina, "No depreciable", [0.0, 0.0, 500.0])
    _escribir(mina, "Edificaciones y construcciones", [0.0, 0.0, 400.0])
    _escribir(libro["Relavera B4"], "Edificaciones y construcciones", [300.0, 0.0, 0.0])
    libro.save(ruta)
    return ruta


def _caso() -> Caso:
    horizonte = Horizonte(primer_ano=2027, anos=ANOS)
    ceros = horizonte.ceros()
    return Caso(
        nombre="Caso de prueba",
        horizonte=horizonte,
        unidades=(
            UnidadProductiva(
                nombre="Mina Alfa",
                tipo="mina",
                produccion=ProduccionDeUnidad(
                    mineral_tratado=horizonte.serie((0.0, 1_000.0, 1_000.0), nombre="tratado"),
                    mineral_extraido=horizonte.serie((0.0, 1_000.0, 1_000.0), nombre="extraido"),
                ),
            ),
            UnidadProductiva(
                nombre="Relavera B4",
                tipo="deposito",
                produccion=ProduccionDeUnidad(mineral_tratado=ceros),
            ),
            UnidadProductiva(
                nombre="Refineria",
                tipo="refineria",
                produccion=ProduccionDeUnidad(mineral_tratado=ceros),
            ),
        ),
        terminos=TerminosComerciales(
            precio_metal_refinado=ceros,
            premio_metal_refinado=ceros,
            precio_metal_en_concentrado=ceros,
            factor_metal_pagable=ceros,
        ),
        datos_comunes=DatosComunes(gastos_administrativos=ceros),
    )


def _con_capex(ruta: Path) -> Caso:
    lectura = leer_capex(ruta)
    assert lectura.valida, [str(i) for i in lectura.incidencias]
    return aplicar(_caso(), [b.capex for b in lectura.bloques])


class TestIdaYVuelta:
    def test_la_plantilla_llena_se_lee(self, libro_de_capex: Path) -> None:
        lectura = leer_capex(libro_de_capex)
        assert lectura.valida, [str(i) for i in lectura.incidencias]
        assert [b.hoja for b in lectura.bloques] == list(PESTANAS)
        assert lectura.horizonte is not None
        assert lectura.horizonte.primer_ano == 2027

    def test_las_tres_pestanas_tienen_la_misma_estructura(self, libro_de_capex: Path) -> None:
        libro = load_workbook(libro_de_capex)
        esperadas = [f.etiqueta for f in CON_DATO_DE_CAPEX]
        for nombre in PESTANAS:
            leidas = [
                str(fila[0].value).strip()
                for fila in libro[nombre].iter_rows(min_row=5, max_col=2)
                if fila[1].value and fila[0].value
            ]
            assert leidas == esperadas

    def test_computo_no_se_consolida_con_maquinaria(self, libro_de_capex: Path) -> None:
        # El libro los junta bajo el codigo MAQ y la plataforma no: un proyecto
        # nuevo puede traer componentes que hoy no existen, y una depreciacion
        # que llega sumada no se puede volver a separar.
        capital = _con_capex(libro_de_capex).unidades[0].capital
        assert capital is not None
        assert capital.por_naturaleza["maquinaria"] == (1_000_000.0, 0.0, 0.0)
        assert capital.por_naturaleza["equipos_de_computo"] == (0.0, 200_000.0, 0.0)

    def test_una_unidad_sin_capital_no_declara_ninguno(self, libro_de_capex: Path) -> None:
        # No es lo mismo que declararlo en ceros: el desglose por mina se
        # llenaria de series nulas.
        assert _con_capex(libro_de_capex).unidades[2].capital is None

    def test_el_caso_leido_calcula(self, libro_de_capex: Path) -> None:
        corrida = calcular(_con_capex(libro_de_capex), MAESTROS)
        assert corrida.capex == (1_300_000.0, 200_000.0, 900_000.0)


class TestConversionDeEscalas:
    def test_los_miles_de_dolares_se_convierten(self, libro_de_capex: Path) -> None:
        # Regla 003: el libro lleva el capital en miles de dolares y alterna las
        # dos escrituras, `$k` y `k$`, dentro de la misma columna.
        capital = _con_capex(libro_de_capex).unidades[1].capital
        assert capital is not None
        assert capital.por_naturaleza["edificaciones"] == (300_000.0, 0.0, 0.0)


class TestLaEtapaSeDeriva:
    def test_el_cierre_es_exactamente_lo_no_depreciable(self, libro_de_capex: Path) -> None:
        capital = _con_capex(libro_de_capex).unidades[0].capital
        assert capital is not None
        assert capital.por_etapa["cierre"] == (0.0, 0.0, 500_000.0)

    def test_lo_anterior_al_primer_ano_con_produccion_es_inicial(
        self, libro_de_capex: Path
    ) -> None:
        capital = _con_capex(libro_de_capex).unidades[0].capital
        assert capital is not None
        assert capital.por_etapa["inicial"] == (1_000_000.0, 0.0, 0.0)
        assert capital.por_etapa["sostenimiento"] == (0.0, 200_000.0, 400_000.0)

    def test_una_unidad_que_no_produce_lo_lleva_todo_a_inicial(self, libro_de_capex: Path) -> None:
        # Sin produccion no hay primer ano de produccion, de modo que no hay
        # frontera entre inicial y sostenimiento.
        capital = _con_capex(libro_de_capex).unidades[1].capital
        assert capital is not None
        assert capital.por_etapa["inicial"] == (300_000.0, 0.0, 0.0)
        assert not any(capital.por_etapa["sostenimiento"])

    def test_las_dos_clasificaciones_cuadran(self, libro_de_capex: Path) -> None:
        for unidad in _con_capex(libro_de_capex).unidades:
            if unidad.capital is None:
                continue
            assert unidad.capital.total_por_etapa() == pytest.approx(
                unidad.capital.total_por_naturaleza()
            )


class TestElDeposito:
    def test_atraviesa_la_cadena_sin_producir(self, libro_de_capex: Path) -> None:
        corrida = calcular(_con_capex(libro_de_capex), MAESTROS)
        assert "Relavera B4" in corrida.depreciacion_tributaria_por_mina

    def test_su_capital_se_deprecia_aunque_no_produzca(self, libro_de_capex: Path) -> None:
        # La regla 013 no deja correr la depreciacion antes de producir, y una
        # relavera de deposito no produce nunca: aplicarsela le anularia el
        # escudo fiscal entero en vez de retrasarlo.
        corrida = calcular(_con_capex(libro_de_capex), MAESTROS)
        assert any(corrida.depreciacion_tributaria_por_mina["Relavera B4"])

    def test_no_entra_como_componente_de_la_refineria(self, libro_de_capex: Path) -> None:
        caso = _con_capex(libro_de_capex)
        assert [u.nombre for u in caso.unidades_mineras] == ["Mina Alfa"]


class TestElAjusteDeCapex:
    def test_sin_ajuste_nada_se_mueve(self, libro_de_capex: Path) -> None:
        caso = _con_capex(libro_de_capex)
        assert calcular(caso, MAESTROS).capex == (1_300_000.0, 200_000.0, 900_000.0)

    def test_el_ajuste_escala_el_capital_y_su_depreciacion(self, libro_de_capex: Path) -> None:
        caso = _con_capex(libro_de_capex)
        sin_ajuste = calcular(caso, MAESTROS)
        con_ajuste = calcular(
            replace(
                caso,
                datos_comunes=replace(
                    caso.datos_comunes,
                    ajuste_de_capex=caso.horizonte.serie((0.5,) * ANOS, nombre="ajuste"),
                ),
            ),
            MAESTROS,
        )
        assert con_ajuste.capex == pytest.approx(tuple(v * 1.5 for v in sin_ajuste.capex))
        assert con_ajuste.depreciacion_tributaria_por_mina["Mina Alfa"] == pytest.approx(
            tuple(v * 1.5 for v in sin_ajuste.depreciacion_tributaria_por_mina["Mina Alfa"])
        )


class TestLasReservas:
    def test_sin_declarar_salen_del_plan_de_la_unidad(self, libro_de_capex: Path) -> None:
        # Es la distincion del libro: la unidad en operacion las trae de su plan
        # de vida de mina y el proyecto las deriva de lo que extrae.
        caso = _con_capex(libro_de_capex)
        assert caso.unidades[0].reservas is None
        corrida = calcular(caso, MAESTROS)
        # Sin reservas declaradas, la unidad se agota exactamente al terminar su
        # plan: la ultima tasa de agotamiento vale uno.
        financiera = corrida.depreciacion_financiera_por_componente["Mina Alfa"]
        assert any(any(serie) for serie in financiera.values())

    def test_declararlas_cambia_el_ritmo(self, libro_de_capex: Path) -> None:
        caso = _con_capex(libro_de_capex)
        sin_declarar = calcular(caso, MAESTROS)
        con_declarar = calcular(
            replace(
                caso,
                unidades=(replace(caso.unidades[0], reservas=100_000.0), *caso.unidades[1:]),
            ),
            MAESTROS,
        )
        # Con mas reservas que lo que el plan extrae, cada ano se agota una
        # fraccion menor del saldo.
        assert (
            con_declarar.depreciacion_financiera_por_mina["Mina Alfa"][2]
            < sin_declarar.depreciacion_financiera_por_mina["Mina Alfa"][2]
        )


class TestIncidencias:
    def test_una_estructura_alterada_se_reporta_y_no_se_adivina(self, libro_de_capex: Path) -> None:
        libro = load_workbook(libro_de_capex)
        libro["Mina Alfa"].cell(row=7, column=1, value="Equipos de oficina")
        libro.save(libro_de_capex)

        lectura = leer_capex(libro_de_capex)
        assert not lectura.valida
        assert any("estructura fija" in i.mensaje for i in lectura.incidencias)

    def test_una_celda_con_texto_se_reporta_con_su_ubicacion(self, libro_de_capex: Path) -> None:
        libro = load_workbook(libro_de_capex)
        fila = _escribir(libro["Mina Alfa"], "No depreciable", [0.0] * ANOS)
        libro["Mina Alfa"].cell(row=fila, column=4, value="quinientos")
        libro.save(libro_de_capex)

        lectura = leer_capex(libro_de_capex)
        assert not lectura.valida
        incidencia = lectura.incidencias[0]
        assert incidencia.hoja == "Mina Alfa"
        assert incidencia.celda == f"D{fila}"

    def test_un_bloque_incompleto_se_reporta(self, libro_de_capex: Path) -> None:
        libro = load_workbook(libro_de_capex)
        hoja = libro["Mina Alfa"]
        hoja.delete_rows(_escribir(hoja, "Edificaciones y construcciones", [0.0] * ANOS))
        libro.save(libro_de_capex)

        assert not leer_capex(libro_de_capex).valida

    def test_importes_sin_concepto_se_reportan(self, libro_de_capex: Path) -> None:
        libro = load_workbook(libro_de_capex)
        hoja = libro["Mina Alfa"]
        fila = _escribir(hoja, "Edificaciones y construcciones", [0.0] * ANOS) + 1
        hoja.cell(row=fila, column=2, value="$k")
        hoja.cell(row=fila, column=3, value=90.0)
        libro.save(libro_de_capex)

        lectura = leer_capex(libro_de_capex)
        assert not lectura.valida
        assert any("sin concepto" in i.mensaje for i in lectura.incidencias)

    def test_un_archivo_que_no_existe_no_revienta(self, tmp_path: Path) -> None:
        lectura = leer_capex(tmp_path / "fantasma.xlsx")
        assert not lectura.valida
        assert lectura.incidencias[0].mensaje == "no existe"


class TestAsociacionPorOrden:
    def test_las_pestanas_se_asocian_por_posicion(self, libro_de_capex: Path) -> None:
        caso = _con_capex(libro_de_capex)
        assert [u.capital is not None for u in caso.unidades] == [True, True, False]

    def test_una_pestana_de_mas_o_de_menos_se_rechaza(self, libro_de_capex: Path) -> None:
        lectura = leer_capex(libro_de_capex)
        with pytest.raises(ErrorDeAsociacion, match="sobra o falta"):
            aplicar(_caso(), [b.capex for b in lectura.bloques][:2])
