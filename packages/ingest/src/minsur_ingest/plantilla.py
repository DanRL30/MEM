"""Lectura de la plantilla canónica: de un libro de Excel a un caso validado.

Es la frontera del motor. Aquí se lee Excel, se traduce el vocabulario del
libro al del catálogo, se validan las celdas y se convierten unidades y escalas.
Del otro lado el motor recibe estructuras limpias y no vuelve a ver una hoja de
cálculo, que es lo que le permite ser puro.

**Aquí y solo aquí se convierten escalas.** El libro alterna dólares y miles de
dólares con factores de 10^3 repartidos por las fórmulas —la regla 003, todavía
sin confirmar— y esa mezcla no se propaga: la plantilla declara sus unidades en
la columna B y el motor trabaja siempre en dólares y toneladas.

**La lectura no se detiene en el primer error.** Acumula incidencias con su hoja
y su celda y las devuelve juntas, porque una plantilla llena a mano llega con
varios problemas a la vez.

La estructura que se lee es la que emite `scripts/generar_plantilla_inputs.py`:
fila 3 con los años desde la columna C, secciones en gris con el nombre de la
unidad, y filas de entrada con concepto en A y unidad de medida en B.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from minsur_engine.capex import CapitalDeUnidad
from minsur_engine.caso import (
    Caso,
    DatosComunes,
    TerminosComerciales,
    UnidadProductiva,
)
from minsur_engine.horizonte import ErrorHorizonte, Horizonte, Serie
from minsur_ingest.incidencias import ErrorDePlantilla, Incidencia
from minsur_ingest.produccion import armar_produccion
from minsur_ingest.sinonimos import canonizar, normalizar

HOJAS_REQUERIDAS = ("Caso", "Opex", "Capex", "Precios")
"""La produccion no esta aqui: viene en una pestana por proyecto.

Sus nombres los declara la propia hoja `Caso`, de modo que no se pueden fijar en
una constante. `Leeme` se ignora.
"""

PRIMERA_FILA_DE_DATOS = 5
PRIMERA_COLUMNA_DE_ANOS = 3
LIMITE_DE_NOMBRE_DE_HOJA = 31

ETAPAS_DE_CAPEX = {
    "capex inicial": "inicial",
    "sostenimiento": "sostenimiento",
    "cierre de mina": "cierre",
    "otros": "otros",
}
NATURALEZAS_DE_CAPEX = {
    "no depreciable": "no_depreciable",
    "maquinaria equipos y vehiculos": "maquinaria",
    "instalaciones y equipos diversos y de comunicaciones": "instalaciones",
    "edificaciones y construcciones": "edificaciones",
}

# Unidad de medida declarada en la columna B, y su factor a la unidad del motor.
# La clave conserva el simbolo de moneda y el de porcentaje: son lo unico que
# distingue "US$" de "miles de US$", y perderlos al normalizar deja pasar un
# factor de mil sin avisar.
FACTORES_DE_ESCALA = {
    "us$": 1.0,
    "us$/ano": 1.0,
    "miles de us$": 1_000.0,
    "mus$": 1_000.0,
    "t": 1.0,
    "tmf": 1.0,
    "kt": 1_000.0,
    "%": 0.01,
    "fraccion": 1.0,
    "dias": 1.0,
    "us$/t": 1.0,
}


def _clave_de_medida(texto: str) -> str:
    """Normaliza una unidad de medida sin borrar su moneda ni su porcentaje."""
    sin_acentos = (
        unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")
    )
    return re.sub(r"\s+", " ", sin_acentos.strip().lower())


@dataclass(frozen=True)
class Lectura:
    """Resultado de leer una plantilla."""

    caso: Caso | None
    incidencias: tuple[Incidencia, ...]

    @property
    def valida(self) -> bool:
        return self.caso is not None


@dataclass(frozen=True)
class _UnidadDeclarada:
    """Una fila de la tabla de unidades de la hoja `Caso`.

    El origen y las etapas viajan aqui y no en la hoja de instrucciones porque
    son los que deciden que filas tiene la pestana de la unidad: sin ellos, de
    una plantilla llena no se puede regenerar la misma plantilla.
    """

    nombre: str
    tipo: str
    metales: str
    origen: str = "yacimiento"
    etapas: str = ""
    entrega_a: str = ""

    @property
    def hoja(self) -> str:
        """Nombre de su pestana, con el limite que impone Excel."""
        limpio = "".join(c for c in self.nombre if c not in r"[]:*?/\'")
        return limpio[:LIMITE_DE_NOMBRE_DE_HOJA] or "Unidad"


@dataclass
class _Cabecera:
    """Lo que la hoja `Caso` declara sobre el caso completo."""

    nombre: str
    escenario: str
    horizonte: Horizonte | None
    unidades: list[_UnidadDeclarada]
    comunes: dict[str, float]


@dataclass
class _Fila:
    """Una fila de datos de la plantilla, ya localizada y traducida."""

    hoja: str
    numero: int
    seccion: str
    subseccion: str
    concepto: str
    metal: str | None
    medida: str
    valores: tuple[object, ...]


def leer_plantilla(ruta: Path, *, escenario: str | None = None) -> Lectura:
    """Lee una plantilla llena y devuelve el caso, o las incidencias que lo impiden.

    `escenario` elige el juego de precios; si no se indica, se usa el que la
    propia plantilla declara en su hoja `Caso`.
    """
    incidencias: list[Incidencia] = []
    if not ruta.exists():
        return Lectura(None, (Incidencia("(archivo)", str(ruta), "no existe"),))

    libro = load_workbook(ruta, data_only=True, read_only=True)
    faltantes = [h for h in HOJAS_REQUERIDAS if h not in libro.sheetnames]
    if faltantes:
        return Lectura(
            None,
            (
                Incidencia(
                    "(libro)",
                    "-",
                    f"faltan las hojas {', '.join(faltantes)}. La plantilla no es la canonica.",
                ),
            ),
        )

    cabecera = _leer_hoja_caso(libro["Caso"], incidencias)
    horizonte = cabecera.horizonte
    if horizonte is None:
        return Lectura(None, tuple(incidencias))

    nombres = [u.nombre for u in cabecera.unidades]
    produccion: dict[str, list[_Fila]] = {}
    for declarada in cabecera.unidades:
        if declarada.hoja not in libro.sheetnames:
            incidencias.append(
                Incidencia(
                    "(libro)",
                    declarada.hoja,
                    f"la unidad {declarada.nombre!r} esta declarada en la hoja Caso y no tiene "
                    "pestana de produccion.",
                )
            )
            continue
        produccion[declarada.nombre] = _leer_pestana_de_unidad(libro[declarada.hoja], horizonte)

    opex = _leer_hoja_de_series(libro["Opex"], horizonte, nombres, incidencias)
    capex = _leer_hoja_de_series(libro["Capex"], horizonte, nombres, incidencias)
    precios = _leer_hoja_de_series(libro["Precios"], horizonte, [], incidencias)
    libro.close()

    unidades = _armar_unidades(cabecera, horizonte, produccion, opex, capex, incidencias)
    terminos = _armar_terminos(precios, horizonte, escenario or cabecera.escenario, incidencias)
    comunes = _armar_datos_comunes(cabecera, horizonte)

    if incidencias or not unidades:
        if not unidades and not incidencias:
            incidencias.append(
                Incidencia("Caso", "A16", "la plantilla no declara unidades productivas")
            )
        return Lectura(None, tuple(incidencias))

    try:
        caso = Caso(
            nombre=cabecera.nombre or ruta.stem,
            horizonte=horizonte,
            unidades=tuple(unidades),
            terminos=terminos,
            datos_comunes=comunes,
        )
    except ValueError as error:
        return Lectura(None, (Incidencia("Caso", "-", str(error)),))

    return Lectura(caso, ())


def leer_o_fallar(ruta: Path, *, escenario: str | None = None) -> Caso:
    """Como `leer_plantilla`, pero levanta con todas las incidencias juntas."""
    lectura = leer_plantilla(ruta, escenario=escenario)
    if lectura.caso is None:
        raise ErrorDePlantilla(lectura.incidencias)
    return lectura.caso


# --- Lectura de hojas ---------------------------------------------------------


def _leer_hoja_caso(hoja: Worksheet, incidencias: list[Incidencia]) -> _Cabecera:
    """Lee la identificación, las unidades declaradas y los datos comunes."""
    etiquetas: dict[str, object] = {}
    unidades: list[_UnidadDeclarada] = []
    comunes: dict[str, float] = {}
    en_unidades = False
    en_comunes = False

    for fila in hoja.iter_rows(min_row=1, max_col=6):
        a, b, c, d, e, f = (celda.value for celda in fila)
        texto = normalizar(str(a)) if a is not None else ""
        if not texto:
            continue
        if texto == "unidades productivas declaradas":
            en_unidades, en_comunes = True, False
            continue
        if texto == "datos comunes al caso":
            en_unidades, en_comunes = False, True
            continue

        if en_unidades:
            if texto in ("unidad",) or b is None:
                continue
            unidades.append(
                _UnidadDeclarada(
                    nombre=str(a).strip(),
                    tipo=str(b).strip(),
                    metales=str(c or ""),
                    origen=str(d).strip() if d else "yacimiento",
                    etapas=str(e or "").strip(),
                    entrega_a=str(f or "").strip(),
                )
            )
        elif en_comunes:
            comunes[texto] = _numero(c, hoja.title, fila[2].coordinate, incidencias) or 0.0
        else:
            etiquetas[texto] = b

    nombre = etiquetas.get("nombre del caso")
    escenario = etiquetas.get("escenario de precios")
    return _Cabecera(
        nombre=str(nombre).strip() if nombre else "",
        escenario=str(escenario).strip() if escenario else "Base",
        horizonte=_armar_horizonte(etiquetas, hoja.title, incidencias),
        unidades=unidades,
        comunes=comunes,
    )


def _armar_horizonte(
    etiquetas: dict[str, object], hoja: str, incidencias: list[Incidencia]
) -> Horizonte | None:
    primer_ano = etiquetas.get("primer ano del horizonte")
    anos = etiquetas.get("numero de anos")
    if not isinstance(primer_ano, int | float) or not isinstance(anos, int | float):
        incidencias.append(
            Incidencia(
                hoja,
                "B6:B7",
                "el primer ano y el numero de anos son obligatorios y deben ser numeros",
            )
        )
        return None
    try:
        return Horizonte(primer_ano=int(primer_ano), anos=int(anos))
    except ErrorHorizonte as error:
        incidencias.append(Incidencia(hoja, "B6:B7", str(error)))
        return None


def _leer_pestana_de_unidad(hoja: Worksheet, horizonte: Horizonte) -> list[_Fila]:
    """Filas de la pestaña de una unidad, con su etiqueta tal cual.

    A diferencia de `_leer_hoja_de_series`, aquí no se separa el metal ni se
    borran los nombres de unidad: en `concentrado alimentado desde San Rafael`
    el nombre es el dato, y quitarlo deja una etiqueta que no significa nada.
    Las filas de sección —`Mina`, `Planta`, `Complejo`— son rótulos y se saltan.
    """
    filas: list[_Fila] = []
    ultima_columna = PRIMERA_COLUMNA_DE_ANOS + horizonte.anos - 1
    for fila in hoja.iter_rows(min_row=PRIMERA_FILA_DE_DATOS, max_col=ultima_columna):
        etiqueta = fila[0].value
        medida = fila[1].value
        if etiqueta is None or not str(etiqueta).strip() or medida is None:
            continue
        filas.append(
            _Fila(
                hoja=hoja.title,
                numero=int(fila[0].row or 0),
                seccion=hoja.title,
                subseccion="",
                concepto=normalizar(str(etiqueta)),
                metal=None,
                medida=_clave_de_medida(str(medida)),
                valores=tuple(celda.value for celda in fila[2:]),
            )
        )
    return filas


def _leer_hoja_de_series(
    hoja: Worksheet, horizonte: Horizonte, unidades: Sequence[str], incidencias: list[Incidencia]
) -> list[_Fila]:
    """Recorre una hoja de años y devuelve sus filas de datos ya traducidas.

    Las filas de sección —las que traen concepto y no traen unidad de medida—
    marcan a quién pertenece lo que viene debajo.
    """
    filas: list[_Fila] = []
    seccion = ""
    subseccion = ""
    ultima_columna = PRIMERA_COLUMNA_DE_ANOS + horizonte.anos - 1

    for fila in hoja.iter_rows(min_row=PRIMERA_FILA_DE_DATOS, max_col=ultima_columna):
        etiqueta = fila[0].value
        medida = fila[1].value
        if etiqueta is None or not str(etiqueta).strip():
            continue
        if medida is None:
            # La hoja de capital anida: la unidad abre seccion y cada etapa
            # abre subseccion. Distinguirlas por el nombre y no por la sangria
            # evita que un espacio de mas desarme la lectura.
            texto = str(etiqueta).strip()
            if normalizar(texto) in ETAPAS_DE_CAPEX:
                subseccion = texto
            else:
                seccion, subseccion = texto, ""
            continue

        concepto, metal = canonizar(str(etiqueta), unidades)
        filas.append(
            _Fila(
                hoja=hoja.title,
                numero=int(fila[0].row or 0),
                seccion=seccion,
                subseccion=subseccion,
                concepto=concepto,
                metal=metal,
                medida=_clave_de_medida(str(medida)),
                valores=tuple(celda.value for celda in fila[2:]),
            )
        )
    return filas


# --- Armado del caso ----------------------------------------------------------


def _armar_unidades(
    cabecera: _Cabecera,
    horizonte: Horizonte,
    produccion: dict[str, list[_Fila]],
    opex: list[_Fila],
    capex: list[_Fila],
    incidencias: list[Incidencia],
) -> list[UnidadProductiva]:
    unidades: list[UnidadProductiva] = []
    for declarada in cabecera.unidades:
        nombre = declarada.nombre
        filas = [
            (fila.concepto, _serie(fila, horizonte, incidencias))
            for fila in produccion.get(nombre, [])
        ]
        costos = {
            str(fila.concepto): _serie(fila, horizonte, incidencias)
            for fila in opex
            if _es_de(fila.seccion, nombre)
        }
        capital = _armar_capital(nombre, horizonte, capex, incidencias)

        unidades.append(
            UnidadProductiva(
                nombre=nombre,
                tipo=declarada.tipo,
                produccion=armar_produccion(
                    nombre,
                    filas,
                    horizonte,
                    incidencias,
                    hoja=declarada.hoja,
                    declaradas=[u.nombre for u in cabecera.unidades],
                ),
                costos={c: s for c, s in costos.items() if any(s)},
                capital=capital,
                origen=declarada.origen,
                etapas=_etapas(declarada),
                entrega_a=declarada.entrega_a or None,
            )
        )
    return unidades


def _etapas(declarada: _UnidadDeclarada) -> tuple[str, ...]:
    """Etapas de la planta, tal como la hoja `Caso` las declara."""
    return tuple(e.strip() for e in declarada.etapas.split(",") if e.strip())


def _armar_capital(
    unidad: str, horizonte: Horizonte, capex: list[_Fila], incidencias: list[Incidencia]
) -> CapitalDeUnidad | None:
    """Reconstruye la doble clasificación del capital de una unidad.

    La hoja anida naturaleza dentro de etapa, así que el mismo importe alimenta
    las dos clasificaciones. Si no cuadraran, `CapitalDeUnidad` lo rechaza: es
    la fila `Check` del libro, convertida en condición.
    """
    por_etapa: dict[str, list[float]] = {}
    por_naturaleza: dict[str, list[float]] = {}

    for fila in capex:
        if not _es_de(fila.seccion, unidad):
            continue
        etapa_actual = ETAPAS_DE_CAPEX.get(normalizar(fila.subseccion))
        naturaleza = NATURALEZAS_DE_CAPEX.get(normalizar(str(fila.concepto)))
        if naturaleza is None or etapa_actual is None:
            continue
        valores = _serie(fila, horizonte, incidencias)
        acumulado_etapa = por_etapa.setdefault(etapa_actual, [0.0] * horizonte.anos)
        acumulado_naturaleza = por_naturaleza.setdefault(naturaleza, [0.0] * horizonte.anos)
        for i, valor in enumerate(valores):
            acumulado_etapa[i] += valor
            acumulado_naturaleza[i] += valor

    if not por_etapa:
        return None
    return CapitalDeUnidad(
        unidad=unidad,
        por_etapa={etapa: tuple(serie) for etapa, serie in por_etapa.items()},
        por_naturaleza={n: tuple(serie) for n, serie in por_naturaleza.items()},
    )


def _armar_terminos(
    precios: list[_Fila],
    horizonte: Horizonte,
    escenario: str,
    incidencias: list[Incidencia],
) -> TerminosComerciales:
    """Toma el juego de precios del escenario elegido y sus condiciones."""
    buscado = normalizar(f"precio escenario {escenario}")
    precio = horizonte.ceros()
    premio = horizonte.ceros()
    factor = horizonte.ceros()

    for fila in precios:
        concepto = normalizar(fila.concepto)
        if concepto == buscado:
            precio = _serie(fila, horizonte, incidencias)
        elif concepto == "premio del metal refinado":
            premio = _serie(fila, horizonte, incidencias)
        elif concepto == "factor de metal pagable":
            factor = _serie(fila, horizonte, incidencias)

    return TerminosComerciales(
        precio_metal_refinado=precio,
        premio_metal_refinado=premio,
        precio_metal_en_concentrado=precio,
        factor_metal_pagable=factor,
    )


def _armar_datos_comunes(cabecera: _Cabecera, horizonte: Horizonte) -> DatosComunes:
    """Convierte los datos comunes en series del horizonte.

    La plantilla los recoge como un valor único por concepto, así que se
    reparten iguales sobre todos los años. Un caso que necesite variarlos año a
    año lo dirá cuando Finanzas revise la plantilla.
    """
    comunes = cabecera.comunes

    def constante(clave: str) -> Serie:
        valor = float(comunes.get(clave, 0.0) or 0.0)
        return tuple(valor for _ in range(horizonte.anos))

    dias = float(comunes.get("dias de working capital", 0.0) or 0.0)
    return DatosComunes(
        gastos_administrativos=constante("gastos administrativos"),
        gestion_social=constante("inversion social"),
        otros_gastos=constante("otros gastos operativos"),
        estudios=constante("estudios"),
        exploraciones=constante("exploraciones no atribuibles a una unidad"),
        predios=constante("servidumbres y usufructos"),
        dias_por_cobrar=tuple(dias for _ in range(horizonte.anos)),
        dias_por_pagar=tuple(dias for _ in range(horizonte.anos)),
        saldo_inicial_de_perdidas=float(
            comunes.get("perdidas tributarias arrastradas", 0.0) or 0.0
        ),
    )


# --- Celdas -------------------------------------------------------------------


def _serie(fila: _Fila, horizonte: Horizonte, incidencias: list[Incidencia]) -> Serie:
    """Convierte los valores de una fila en una serie del horizonte."""
    factor = FACTORES_DE_ESCALA.get(fila.medida, 1.0)
    valores: list[float] = []
    for i in range(horizonte.anos):
        bruto = fila.valores[i] if i < len(fila.valores) else None
        celda = f"{get_column_letter(PRIMERA_COLUMNA_DE_ANOS + i)}{fila.numero}"
        numero = _numero(bruto, fila.hoja, celda, incidencias)
        valores.append((numero or 0.0) * factor)
    return tuple(valores)


def _numero(valor: object, hoja: str, celda: str, incidencias: list[Incidencia]) -> float | None:
    """Un valor de celda como número, o una incidencia localizada.

    Una celda vacía es cero y no es un error: significa que el concepto no
    aplica ese año, que es información legítima y frecuente.
    """
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        incidencias.append(Incidencia(hoja, celda, f"se esperaba un numero y hay {valor!r}"))
        return None
    if isinstance(valor, int | float):
        return float(valor)
    try:
        return float(str(valor).replace(",", "").replace("%", "").strip())
    except ValueError:
        incidencias.append(
            Incidencia(hoja, celda, f"se esperaba un numero y hay {str(valor)[:40]!r}")
        )
        return None


def _es_de(seccion: str, unidad: str) -> bool:
    """Si una sección pertenece a una unidad.

    La plantilla escribe la sección como `San Rafael (mina)`, y el libro a
    veces solo el nombre. Comparar por prefijo normalizado cubre ambas.
    """
    normalizada = normalizar(seccion)
    objetivo = normalizar(unidad)
    return normalizada == objetivo or normalizada.startswith(f"{objetivo} ")
